from __future__ import annotations

import json
import os
import shlex
import shutil
import subprocess
import tempfile
import traceback
from concurrent.futures import Future, ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from threading import Lock
from typing import Dict, List, Optional, Tuple
from uuid import uuid4

import pandas as pd

from .ids import build_run_id
from .models import WorkflowStepResult
from .state import (
    record_checkpoint_metadata,
    record_step,
    update_artifacts,
    update_background_task,
    update_context,
)


ALL_ANALYSIS_TARGETS = [
    "analyze.stage.hierarchical",
    "analyze.stage.polypharmacology",
    "analyze.stage.rmsd",
    "analyze.stage.reports",
    "analyze.stage.visualizations",
    "analyze.stage.structure_quality",
    "analyze.interactions.pandamap",
    "analyze.interactions.prolif",
    "analyze.interactions.ligplot",
    "analyze.interactions.poseview",
    "analyze.interactions.clean",
    "analyze.visuals.py3dmol",
    "analyze.visuals.pymol",
]
_INTERACTION_TARGETS_ROUTED_TO_CLEAN = {
    "analyze.interactions.pandamap",
    "analyze.interactions.prolif",
    "analyze.interactions.ligplot",
}
_BLOCKED_LEGACY_STAGE_TARGETS = {
    "analyze.stage.structure_quality",
    "analyze.visuals.pymol",
}

DEFAULT_CLEAN_INTERACTION_DATASET_ROOT = Path(
    os.environ.get("DOCKFORGE_CLEAN_INTERACTION_DATASET_ROOT", str(Path.cwd()))
).expanduser().resolve()
_ANALYSIS_CONFIG_CANDIDATES = (
    "post_docking_analysis.yaml",
    "post_docking_analysis.yml",
    "post_docking_analysis.json",
    "analysis_config.yaml",
    "analysis_config.yml",
    "analysis_config.json",
)


def _config_by_path(config: Dict[str, object], path: str, default=None):
    current = config
    for token in str(path).split("."):
        if not isinstance(current, dict) or token not in current:
            return default
        current = current[token]
    return current


def _discover_analysis_config_file(project_dir: Optional[str]) -> Optional[Path]:
    if not project_dir:
        return None
    root = Path(project_dir).expanduser().resolve()
    candidates: List[Path] = []
    for name in _ANALYSIS_CONFIG_CANDIDATES:
        candidates.append(root / name)
        candidates.append(root / "config" / name)
    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            return candidate
    return None


def _resolve_analysis_config_overrides(
    *,
    project_dir: Optional[str],
    config_file: Optional[str],
) -> Tuple[str, Dict[str, object], List[str]]:
    notes: List[str] = []
    resolved_path = ""
    selected: Optional[Path] = Path(config_file).expanduser().resolve() if config_file else _discover_analysis_config_file(project_dir)
    if selected is None:
        return resolved_path, {}, notes
    try:
        from post_docking_analysis.config_manager import load_config_overrides

        overrides = load_config_overrides(str(selected))
        resolved_path = str(selected)
        notes.append(f"Loaded analysis config overrides: {selected}")
        return resolved_path, overrides, notes
    except Exception as exc:
        notes.append(f"Could not load analysis config overrides from {selected}: {exc}")
        return str(selected), {}, notes


def _raise_legacy_target_migration_error(target: str) -> None:
    migration_hints = {
        "analyze.stage.structure_quality": (
            "Use `python main.py analyze favorite-engine --project-dir <project> --favorite-engine <engine> --analysis-scope full` "
            "or `python main.py analyze comparative --project-dir <project> --analysis-scope full`."
        ),
        "analyze.visuals.pymol": (
            "Use `python main.py analyze visuals py3dmol --project-dir <project>` for supported 3D visuals."
        ),
    }
    hint = migration_hints.get(
        target,
        "Use `python main.py analyze clean --project-dir <project>` for the hardened post-docking interaction contract.",
    )
    raise ValueError(
        f"Target {target} is retired because it depended on legacy PostDockingAnalysisPipeline routes. {hint}"
    )


class BackgroundTaskManager:
    """Simple in-process task runner for non-blocking interactive actions."""

    def __init__(self, max_workers: int = 2):
        self._executor = ThreadPoolExecutor(max_workers=max_workers, thread_name_prefix="dockforge-task")
        self._lock = Lock()
        self._futures: Dict[str, Future] = {}

    def submit(
        self,
        *,
        root: Path,
        label: str,
        target: str,
        fn,
        **kwargs,
    ) -> str:
        task_id = build_run_id("task", suffix=uuid4().hex[:8])
        root_path = Path(root).expanduser().resolve()
        update_background_task(
            root_path,
            task_id,
            label=label,
            target=target,
            status="queued",
            progress=0,
            note="Queued",
        )
        future = self._executor.submit(
            self._run_wrapped,
            task_id=task_id,
            root=root_path,
            label=label,
            target=target,
            fn=fn,
            kwargs=kwargs,
        )
        with self._lock:
            self._futures[task_id] = future
        return task_id

    def _run_wrapped(
        self,
        *,
        task_id: str,
        root: Path,
        label: str,
        target: str,
        fn,
        kwargs: Dict[str, object],
    ) -> None:
        update_background_task(
            root,
            task_id,
            label=label,
            target=target,
            status="running",
            progress=5,
            note="Running",
        )
        try:
            result = fn(**kwargs)
            result_status = str(getattr(result, "status", "completed"))
            mapped_status = "completed" if result_status in {"completed", "validated", "partial"} else "failed"
            update_background_task(
                root,
                task_id,
                label=label,
                target=target,
                status=mapped_status,
                progress=100,
                note=f"Finished with status: {result_status}",
                result_status=result_status,
            )
        except Exception as exc:
            update_background_task(
                root,
                task_id,
                label=label,
                target=target,
                status="failed",
                progress=100,
                note=f"Task failed: {exc}",
                error=traceback.format_exc(limit=5),
            )
        finally:
            with self._lock:
                self._futures.pop(task_id, None)


def _result(
    target: str,
    status: str,
    root: Path,
    inputs: Optional[Dict[str, object]] = None,
    outputs: Optional[Dict[str, object]] = None,
    notes: Optional[List[str]] = None,
) -> WorkflowStepResult:
    root = Path(root).expanduser().resolve()
    record_step(root, target, status, inputs=inputs, outputs=outputs, notes=notes)
    return WorkflowStepResult(
        target=target,
        status=status,
        root=str(root),
        inputs=inputs or {},
        outputs=outputs or {},
        notes=notes or [],
    )


def _safe_root(project_dir: Optional[str] = None, output_dir: Optional[str] = None) -> Path:
    if project_dir:
        return Path(project_dir).expanduser().resolve()
    if output_dir:
        return Path(output_dir).expanduser().resolve()
    return Path.cwd().resolve()


def _prune_empty_dirs(root: Path) -> int:
    removed = 0
    root = Path(root).expanduser().resolve()
    directories = sorted(
        (path for path in root.rglob("*") if path.is_dir()),
        key=lambda path: len(path.parts),
        reverse=True,
    )
    for directory in directories:
        try:
            if directory == root:
                continue
            if any(directory.iterdir()):
                continue
            directory.rmdir()
            removed += 1
        except Exception:
            continue
    return removed


def is_canonical_project(project_dir: str) -> bool:
    from docking.project_layout import manifest_path

    root = Path(project_dir).expanduser().resolve()
    if manifest_path(root).exists():
        return True
    # Backward compatibility: older runs may keep the manifest at project root
    # even when the docking_legacy layout exists.
    return (root / "project_manifest.json").exists() or (root / "4-Docking" / "project_manifest.json").exists()


def _analysis_scope_for_stage_target(target: str) -> str:
    """
    Map legacy stage targets onto unified analysis scopes.
    """
    if target == "analyze.stage.reports":
        return "report_only"
    if target in {
        "analyze.stage.hierarchical",
        "analyze.stage.polypharmacology",
        "analyze.stage.rmsd",
    }:
        return "comparison_only"
    return "full"


def _has_existing_legacy_gnina_layout(detected: Dict[str, object]) -> bool:
    """Require real GNINA output signals before treating a root as legacy."""
    sdf_folder = detected.get("sdf_folder")
    if not isinstance(sdf_folder, Path):
        return False
    return sdf_folder.exists()


def _infer_existing_legacy_engines(root: Path, detected: Dict[str, object]) -> List[str]:
    """Infer active engines from an already-populated legacy docking tree."""
    dock_root = root / "4-Docking" if (root / "4-Docking").is_dir() else root
    engines: List[str] = []

    if detected.get("sdf_folder") or (dock_root / "gnina_out").is_dir():
        engines.append("gnina")
    if (dock_root / "vina_out").is_dir():
        engines.append("vina")
    if (dock_root / "smina_out").is_dir():
        engines.append("smina")
    if (dock_root / "autodock4_out").is_dir():
        engines.append("autodock4")

    return engines or ["gnina"]


def _update_project_context_from_manifest(root: Path) -> None:
    from docking.project_layout import load_manifest, manifest_path

    manifest = load_manifest(root)
    update_context(
        root,
        project_root=str(Path(root).expanduser().resolve()),
        selected_engines=manifest.get("engines", []),
        enabled_panels=manifest.get("enabled_panels", []),
        favorite_engine=manifest.get("favorite_engine", "") or "",
        pair_mode=manifest.get("pairlist_mode", "") or "",
        layout_profile=manifest.get("layout_profile", "") or "",
        raw_proteins_dir=manifest.get("raw_proteins_dir", "") or "",
        raw_ligands_dir=manifest.get("raw_ligands_dir", "") or "",
        raw_ligands_sdf_dir=manifest.get("raw_ligands_sdf_dir", "") or "",
        prepared_proteins_dir=manifest.get("prepared_proteins_dir", "") or "",
        prepared_ligands_dir=manifest.get("prepared_ligands_dir", "") or "",
        docking_root=manifest.get("docking_root", "") or "",
        post_docking_root=manifest.get("post_docking_root", "") or "",
    )
    update_artifacts(
        root,
        project_manifest=str(manifest_path(root)),
        pairlist_file=manifest.get("pairlist_file", "") or "",
        raw_proteins_dir=manifest.get("raw_proteins_dir", "") or "",
        raw_ligands_dir=manifest.get("raw_ligands_dir", "") or "",
        raw_ligands_sdf_dir=manifest.get("raw_ligands_sdf_dir", "") or "",
        prepared_proteins_dir=manifest.get("prepared_proteins_dir", "") or "",
        prepared_ligands_dir=manifest.get("prepared_ligands_dir", "") or "",
        docking_root=manifest.get("docking_root", "") or "",
        post_docking_root=manifest.get("post_docking_root", "") or "",
    )


def run_workflow_init(
    project_dir: str,
    engines: Optional[List[str]] = None,
    project_name: str = "",
    favorite_engine: str = "",
    layout_profile: str = "docking_legacy",
) -> WorkflowStepResult:
    from docking.engine_registry import normalize_engines
    from docking.project_layout import bootstrap_project_layout, manifest_path
    from post_docking_analysis.gnina_hpc_adapter import detect_gnina_layout

    root = Path(project_dir).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    selected_engines = normalize_engines(engines or ["gnina", "vina", "smina", "autodock4"])

    if is_canonical_project(str(root)):
        summary = bootstrap_project_layout(
            root,
            engines=selected_engines,
            project_name=project_name,
            favorite_engine=favorite_engine,
            notes=["workflow_init_completed_missing_structure"],
            layout_profile=layout_profile,
        )
        _update_project_context_from_manifest(root)
        update_artifacts(root, canonical_layout_type=summary["layout_type"])
        return _result(
            "workflow.init_project",
            "completed",
            root,
            inputs={"project_dir": str(root), "engines": selected_engines},
            outputs=summary,
            notes=["Canonical project already existed; ensured required structure is present."],
        )

    detected = detect_gnina_layout(root)
    if _has_existing_legacy_gnina_layout(detected):
        legacy_manifest = manifest_path(root, "docking_legacy")
        root_manifest = manifest_path(root, "canonical")
        if not legacy_manifest.exists() and not root_manifest.exists():
            inferred_engines = _infer_existing_legacy_engines(root, detected)
            summary = bootstrap_project_layout(
                root,
                engines=inferred_engines,
                project_name=project_name,
                favorite_engine=favorite_engine or ("gnina" if "gnina" in inferred_engines else ""),
                notes=[
                    "workflow_init_backfilled_manifest_for_legacy_layout",
                    "Detected existing docking outputs and created a workflow manifest for resume operations.",
                ],
                layout_profile=layout_profile,
            )
            _update_project_context_from_manifest(root)
            update_artifacts(root, canonical_layout_type=summary["layout_type"])
            return _result(
                "workflow.init_project",
                "completed",
                root,
                inputs={"project_dir": str(root), "engines": inferred_engines},
                outputs={
                    **summary,
                    "registered_as": "legacy_gnina_project_with_manifest",
                },
                notes=[
                    "Detected an existing GNINA local/HPC project layout.",
                    "Backfilled a workflow manifest so interactive prep and docking stages can resume safely.",
                ],
            )

        update_context(root, project_root=str(root), selected_engines=["gnina"], favorite_engine=favorite_engine or "gnina")
        update_artifacts(
            root,
            legacy_project_type="gnina_hpc_or_local",
            gnina_sdf_folder=str(detected["sdf_folder"]) if detected.get("sdf_folder") else "",
            gnina_log_folder=str(detected["log_folder"]) if detected.get("log_folder") else "",
            receptors_folder=str(detected["receptors_folder"]) if detected.get("receptors_folder") else "",
            pairlist_file=str(detected["pairlist_file"]) if detected.get("pairlist_file") else "",
        )
        return _result(
            "workflow.init",
            "completed",
            root,
            inputs={"project_dir": str(root), "engines": ["gnina"]},
            outputs={
                "project_root": str(root),
                "layout_type": detected.get("layout", "legacy_gnina"),
                "registered_as": "legacy_gnina_project",
            },
            notes=[
                "Detected an existing GNINA local/HPC project layout.",
                "Workflow state was initialized without forcing canonical multi-engine scaffolding.",
            ],
        )

    summary = bootstrap_project_layout(
        root,
        engines=selected_engines,
        project_name=project_name,
        favorite_engine=favorite_engine,
        notes=[
            "workflow_init_created_canonical_scaffold",
            "Canonical layout remains compatible with multi-engine docking and post-docking analysis.",
        ],
        layout_profile=layout_profile,
    )
    _update_project_context_from_manifest(root)
    update_artifacts(root, canonical_layout_type=summary["layout_type"])
    return _result(
        "workflow.init_project",
        "completed",
        root,
        inputs={"project_dir": str(root), "engines": selected_engines},
        outputs=summary,
        notes=[
            "Created staged docking project scaffold with workflow state, manifest, pairlist stub, and engine directories.",
            "Legacy GNINA HPC projects remain supported separately by layout detection in post-docking analysis.",
        ],
    )


def run_workflow_clone_checkpoint(
    source_project_dir: str,
    target_project_dir: Optional[str] = None,
    layout_profile: Optional[str] = None,
) -> WorkflowStepResult:
    from docking.project_layout import detect_layout_profile, load_manifest

    source_root = Path(source_project_dir).expanduser().resolve()
    if not source_root.exists():
        raise FileNotFoundError(f"Source project root does not exist: {source_root}")

    target_root = (
        Path(target_project_dir).expanduser().resolve()
        if target_project_dir
        else source_root.parent / f"{source_root.name}_checkpoint"
    )

    if target_root.exists() and any(target_root.iterdir()):
        raise FileExistsError(
            f"Target checkpoint workspace already exists and is not empty: {target_root}"
        )
    if not target_root.exists():
        target_root.parent.mkdir(parents=True, exist_ok=True)

    source_layout_profile = "docking_legacy"
    source_engines: Optional[List[str]] = None
    source_favorite_engine = ""
    source_project_name = source_root.name
    try:
        source_layout_profile = detect_layout_profile(source_root)
        if is_canonical_project(str(source_root)):
            source_manifest = load_manifest(source_root)
            manifest_engines = source_manifest.get("engines", []) or []
            source_engines = [str(engine).strip() for engine in manifest_engines if str(engine).strip()]
            source_favorite_engine = str(source_manifest.get("favorite_engine", "") or "")
            source_project_name = str(source_manifest.get("project_name", "") or source_project_name)
    except Exception:
        # Keep cloning resilient even if source manifest metadata is incomplete.
        pass

    shutil.copytree(
        source_root,
        target_root,
        dirs_exist_ok=target_root.exists(),
        ignore=shutil.ignore_patterns(".git", "__pycache__", ".DS_Store"),
    )

    selected_layout_profile = layout_profile or source_layout_profile
    run_workflow_init(
        str(target_root),
        engines=source_engines,
        project_name=source_project_name,
        favorite_engine=source_favorite_engine,
        layout_profile=selected_layout_profile,
    )

    marker = target_root / "CHECKPOINT_REVISE_WORKSPACE.md"
    marker.write_text(
        "\n".join(
            [
                "# Checkpoint And Revise Workspace",
                "",
                f"Source project: {source_root}",
                f"Created at: {datetime.now().isoformat(timespec='seconds')}",
                f"Layout profile: {selected_layout_profile}",
                "",
                "This workspace is intended for checkpoint freeze and iterative revision.",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    checkpoint_record = record_checkpoint_metadata(
        target_root,
        source_project_dir=str(source_root),
        target_project_dir=str(target_root),
        layout_profile=selected_layout_profile,
        marker_file=str(marker),
        note="Created from workflow.clone_checkpoint",
        metadata={
            "source_layout_profile": source_layout_profile,
            "requested_layout_profile": str(layout_profile or ""),
            "resolved_layout_profile": selected_layout_profile,
        },
    )

    return _result(
        "workflow.clone_checkpoint",
        "completed",
        target_root,
        inputs={
            "source_project_dir": str(source_root),
            "requested_target_project_dir": str(target_project_dir or ""),
            "layout_profile": str(layout_profile or ""),
        },
        outputs={
            "source_project_dir": str(source_root),
            "target_project_dir": str(target_root),
            "layout_profile": selected_layout_profile,
            "marker_file": str(marker),
            "checkpoint_id": str(checkpoint_record.get("checkpoint_id", "")),
        },
        notes=["Created a Checkpoint & Revise workspace clone while preserving source manifest defaults."],
    )


def run_workflow_clone_maturation(
    source_project_dir: str,
    target_project_dir: Optional[str] = None,
    layout_profile: Optional[str] = None,
) -> WorkflowStepResult:
    """Backward-compatible alias for older command surfaces."""
    return run_workflow_clone_checkpoint(
        source_project_dir=source_project_dir,
        target_project_dir=target_project_dir,
        layout_profile=layout_profile,
    )


def run_legacy_pdb_interactive(output_dir: Optional[str] = None) -> WorkflowStepResult:
    from interactive_pipeline import run_interactive_pipeline

    root = _safe_root(output_dir=output_dir)
    run_interactive_pipeline(output_dir=output_dir)
    outputs = {"output_dir": str(Path(output_dir).expanduser().resolve())} if output_dir else {}
    update_artifacts(root, legacy_interactive_output=outputs.get("output_dir", ""))
    return _result("pdb.run", "completed", root, outputs=outputs, notes=["Executed legacy interactive PDB workflow."])


def run_pdb_fetch(pdbs: str, output_dir: str) -> WorkflowStepResult:
    from cli_pipeline import parse_pdb_list
    from core_pipeline import MolecularDockingPipeline

    pdb_list = parse_pdb_list(pdbs)
    root = _safe_root(output_dir=output_dir)
    pipeline = MolecularDockingPipeline(output_dir)
    fetched: List[str] = []
    for pdb_id in pdb_list:
        fetched.append(pipeline.fetch_pdb(pdb_id))
    update_artifacts(root, fetched_pdbs=fetched)
    return _result(
        "pdb.fetch",
        "completed",
        root,
        inputs={"pdbs": pdb_list},
        outputs={"files": fetched, "output_dir": str(Path(output_dir).expanduser().resolve())},
    )


def run_pdb_collect(
    project_dir: str,
    pdbs: str,
    selection_mode: str = "heuristic",
    preserve_metals: bool = True,
    preserve_cofactors: bool = True,
    preserve_full_receptor: bool = False,
    preserve_residues: Optional[List[str]] = None,
    preferred_ligands: Optional[List[str]] = None,
) -> WorkflowStepResult:
    from cli_pipeline import EXCEL_AVAILABLE, parse_pdb_list, run_single_pdb_cli
    from interactive_pipeline import run_single_pdb_analysis
    from core_pipeline import MolecularDockingPipeline
    from docking.project_layout import ensure_project_layout, load_manifest

    if not EXCEL_AVAILABLE:
        raise RuntimeError("openpyxl is required for project-aware PDB collection")

    root = Path(project_dir).expanduser().resolve()
    if not is_canonical_project(str(root)):
        run_workflow_init(str(root), layout_profile="docking_legacy")
    layout = ensure_project_layout(root, "docking_legacy")
    pdb_list = parse_pdb_list(pdbs)
    pipeline = MolecularDockingPipeline(str(layout["raw_proteins"]))

    from openpyxl import Workbook, load_workbook

    workbook_path = layout["raw_proteins"] / "multi_pdb_analysis.xlsx"
    workbook = load_workbook(workbook_path) if workbook_path.exists() else Workbook()
    if "Sheet" in workbook.sheetnames:
        workbook.remove(workbook["Sheet"])

    warnings: List[str] = []
    copied_ligands: List[str] = []
    collected: List[str] = []

    collect_config = {
        "ligand_mode": selection_mode if selection_mode in {"auto", "heuristic"} else "auto",
        "preferred_ligands": preferred_ligands or [],
        "cleaning": {
            "default": "common",
            "preserve_metals": preserve_metals,
            "preserve_cofactors": preserve_cofactors,
            "preserve_full_receptor": preserve_full_receptor,
            "preserve_residues": [item.upper() for item in (preserve_residues or []) if item],
        },
    }

    for pdb_id in pdb_list:
        if selection_mode == "interactive":
            success = run_single_pdb_analysis(pipeline, pdb_id, workbook)
        else:
            success = run_single_pdb_cli(pipeline, pdb_id, collect_config, workbook)
        if not success:
            warnings.append(f"Failed to collect {pdb_id}")
            continue
        collected.append(pdb_id)
        for ligand_pdb in sorted(layout["raw_proteins"].glob(f"{pdb_id}_ligand_*.pdb")):
            ligand_copy = layout["raw_ligands"] / ligand_pdb.name
            shutil.copy2(ligand_pdb, ligand_copy)
            copied_ligands.append(str(ligand_copy))
            try:
                ligand_pdb.unlink()
            except OSError as exc:
                warnings.append(f"Failed to remove ligand artifact from raw proteins ({ligand_pdb.name}): {exc}")
            sdf_output = layout["raw_ligands_sdf"] / f"{ligand_pdb.stem}.sdf"
            if shutil.which("obabel"):
                completed = subprocess.run(
                    ["obabel", str(ligand_copy), "-O", str(sdf_output)],
                    capture_output=True,
                    text=True,
                )
                if completed.returncode != 0:
                    warnings.append(f"Open Babel conversion failed for {ligand_pdb.name}: {completed.stderr.strip()}")
            else:
                warnings.append("Open Babel not found; raw ligand SDF conversion skipped.")

    if not workbook.sheetnames:
        ws = workbook.create_sheet("Summary")
        ws.append(["PDB_ID", "Property", "Value"])

    workbook.save(workbook_path)

    manifest = load_manifest(root)
    source_paths = dict(manifest.get("source_paths", {}))
    source_paths.update(
        {
            "raw_proteins": str(layout["raw_proteins"]),
            "raw_ligands": str(layout["raw_ligands"]),
            "raw_ligands_sdf": str(layout["raw_ligands_sdf"]),
            "excel_path": str(workbook_path),
        }
    )
    manifest["source_paths"] = source_paths
    from docking.project_layout import save_manifest

    save_manifest(root, manifest)
    _update_project_context_from_manifest(root)
    update_artifacts(
        root,
        multi_pdb_analysis=str(workbook_path),
        collected_pdbs=collected,
    )
    if warnings:
        warning_file = layout["metadata"] / "pdb_collect_warnings.txt"
        warning_file.write_text("\n".join(sorted(set(warnings))), encoding="utf-8")
    return _result(
        "pdb.collect.append" if len(pdb_list) > 1 else "pdb.collect",
        "completed" if len(collected) == len(pdb_list) else "partial",
        root,
        inputs={"pdbs": pdb_list},
        outputs={
            "raw_proteins_dir": str(layout["raw_proteins"]),
            "raw_ligands_dir": str(layout["raw_ligands"]),
            "raw_ligands_sdf_dir": str(layout["raw_ligands_sdf"]),
            "workbook": str(workbook_path),
            "collected_pdbs": collected,
            "copied_ligands": copied_ligands,
            "selection_mode": selection_mode,
            "preserve_metals": preserve_metals,
            "preserve_cofactors": preserve_cofactors,
            "preserve_full_receptor": preserve_full_receptor,
            "preserve_residues": [item.upper() for item in (preserve_residues or []) if item],
        },
        notes=warnings,
    )


def run_pdb_cli(pdbs: str, output_dir: str, config_file: Optional[str] = None) -> WorkflowStepResult:
    from cli_pipeline import load_config, parse_pdb_list, run_batch_analysis
    from core_pipeline import MolecularDockingPipeline

    pdb_list = parse_pdb_list(pdbs)
    root = _safe_root(output_dir=output_dir)
    pipeline = MolecularDockingPipeline(output_dir)
    config = load_config(config_file) if config_file else {}
    results = run_batch_analysis(pipeline, pdb_list, config, Path(output_dir))
    success = all(results.values()) if results else False
    status = "completed" if success else "partial"
    update_artifacts(root, pdb_output_dir=str(Path(output_dir).expanduser().resolve()))
    return _result(
        "pdb.run",
        status,
        root,
        inputs={"pdbs": pdb_list, "config_file": config_file or ""},
        outputs={"results": results, "output_dir": str(Path(output_dir).expanduser().resolve())},
        notes=[] if success else ["One or more PDB entries failed."],
    )


def run_prepare_pairlist(
    project_dir: str,
    mode: Optional[str],
    prepared_proteins: Optional[str] = None,
    prepared_ligands: Optional[str] = None,
    excel_path: Optional[str] = None,
    default_site_id: str = "site_1",
    default_box_size: float = 20.0,
    curated_receptors: Optional[List[str]] = None,
    curated_ligands: Optional[List[str]] = None,
    curated_mapping: Optional[Dict[str, List[str]]] = None,
    iterative: bool = False,
    round_id: Optional[str] = None,
    freeze: bool = False,
    prompt_protein_aliases: bool = False,
    prompt_ligand_aliases: bool = False,
) -> WorkflowStepResult:
    from docking.preparation.excel_sites import load_site_catalog_from_summary
    from docking.preparation.pair_curation import materialize_pair_curation_round, upsert_pair_curation_round
    from docking.preparation.pairlist_builder import build_pairlists
    from docking.preparation.project_aliases import ensure_project_alias_files, prompt_project_aliases
    from docking.project_layout import ensure_project_layout, load_manifest, pair_curation_state_path, save_manifest

    root = Path(project_dir).expanduser().resolve()
    if not is_canonical_project(str(root)):
        run_workflow_init(str(root), layout_profile="docking_legacy")
    layout = ensure_project_layout(root, "docking_legacy")
    manifest = load_manifest(root)
    resolved_prepared_proteins = Path(
        prepared_proteins or manifest.get("prepared_proteins_dir") or layout["prepared_proteins"]
    ).expanduser().resolve()
    resolved_prepared_ligands = Path(
        prepared_ligands or manifest.get("prepared_ligands_dir") or layout["prepared_ligands"]
    ).expanduser().resolve()
    resolved_excel = Path(
        excel_path or manifest.get("source_paths", {}).get("excel_path") or (layout["raw_proteins"] / "multi_pdb_analysis.xlsx")
    ).expanduser().resolve()
    layout_profile = manifest.get("layout_profile", "docking_legacy") or "docking_legacy"
    if not resolved_excel.exists():
        nested_workbooks = sorted(root.glob("*/2-Raw_Protien/multi_pdb_analysis.xlsx"))
        detail = ""
        if len(nested_workbooks) == 1:
            suggestion = nested_workbooks[0].parent.parent
            detail = (
                f" The selected project root may be a parent directory. "
                f"Try using project root: {suggestion}"
            )
        elif len(nested_workbooks) > 1:
            options = ", ".join(str(path.parent.parent) for path in nested_workbooks[:5])
            detail = (
                f" The selected project root may be a parent directory. "
                f"Candidate project roots under it: {options}"
            )
        raise FileNotFoundError(f"Site workbook not found: {resolved_excel}.{detail}")

    site_catalog = load_site_catalog_from_summary(
        resolved_excel,
        allow_missing_coordinates=True,
    )
    alias_outputs = ensure_project_alias_files(
        project_root=root,
        prepared_proteins=resolved_prepared_proteins,
        prepared_ligands=resolved_prepared_ligands,
        site_catalog=site_catalog,
        layout_profile=layout_profile,
    )
    if prompt_protein_aliases or prompt_ligand_aliases:
        alias_outputs = prompt_project_aliases(
            project_root=root,
            prompt_protein_aliases=prompt_protein_aliases,
            prompt_ligand_aliases=prompt_ligand_aliases,
            layout_profile=layout_profile,
        )

    if iterative or freeze:
        if iterative:
            if not mode:
                raise ValueError("Pair-curation round updates require a pairlist mode")
            summary = upsert_pair_curation_round(
                project_root=root,
                prepared_proteins=resolved_prepared_proteins,
                prepared_ligands=resolved_prepared_ligands,
                excel_path=resolved_excel,
                mode=mode,
                default_site_id=default_site_id,
                default_box_size=default_box_size,
                curated_receptors=curated_receptors,
                curated_ligands=curated_ligands,
                curated_mapping=curated_mapping,
                layout_profile=layout_profile,
                round_id=round_id,
                freeze=freeze,
            )
        else:
            summary = materialize_pair_curation_round(
                root,
                round_id=round_id,
                layout_profile=layout_profile,
            )
            summary["pair_curation_state_file"] = str(pair_curation_state_path(root, layout_profile))
    else:
        if not mode:
            raise ValueError("Pairlist generation requires a pairlist mode")
        summary = build_pairlists(
            project_root=root,
            prepared_proteins=resolved_prepared_proteins,
            prepared_ligands=resolved_prepared_ligands,
            excel_path=resolved_excel,
            mode=mode,
            default_site_id=default_site_id,
            default_box_size=default_box_size,
            curated_receptors=curated_receptors,
            curated_ligands=curated_ligands,
            curated_mapping=curated_mapping,
            layout_profile=layout_profile,
        )

    summary.setdefault("protein_alias_file", alias_outputs.get("protein_alias_file", ""))
    summary.setdefault("ligand_alias_file", alias_outputs.get("ligand_alias_file", ""))

    if summary.get("pairlist_file"):
        manifest["pairlist_file"] = summary["pairlist_file"]
    manifest["pairlist_mode"] = summary.get("pair_mode") or mode or manifest.get("pairlist_mode", "")
    manifest["has_cocrystal_benchmark_rows"] = summary.get("has_cocrystal_benchmark_rows", False)
    source_paths = dict(manifest.get("source_paths", {}))
    source_paths.update(
        {
            "protein_alias_file": summary.get("protein_alias_file") or alias_outputs.get("protein_alias_file", ""),
            "ligand_alias_file": summary.get("ligand_alias_file") or alias_outputs.get("ligand_alias_file", ""),
        }
    )
    manifest["source_paths"] = source_paths
    if summary.get("pair_curation_state_file"):
        manifest["pair_curation_state_file"] = summary["pair_curation_state_file"]
    if summary.get("round_id"):
        manifest["latest_pair_round"] = summary["round_id"]
    save_manifest(root, manifest)

    update_context(root, pair_mode=str(summary.get("pair_mode") or mode or ""))
    update_artifacts(
        root,
        pairlist_file=summary.get("pairlist_file", ""),
        pair_intent_file=summary.get("pair_intent_file", ""),
        pair_curation_state_file=summary.get("pair_curation_state_file", ""),
        protein_alias_file=summary.get("protein_alias_file", "") or alias_outputs.get("protein_alias_file", ""),
        ligand_alias_file=summary.get("ligand_alias_file", "") or alias_outputs.get("ligand_alias_file", ""),
    )
    return _result(
        "prep.pairlist",
        "completed",
        root,
        inputs={
            "mode": mode or "",
            "iterative": iterative,
            "freeze": freeze,
            "round_id": round_id or "",
            "prompt_protein_aliases": prompt_protein_aliases,
            "prompt_ligand_aliases": prompt_ligand_aliases,
        },
        outputs=summary,
        notes=summary.get("warnings", []),
    )


def run_pdb_batch(config_file: str, output_dir: str) -> WorkflowStepResult:
    from batch_pdb_preparation import BatchPDBPreparationPipeline

    root = _safe_root(output_dir=output_dir)
    pipeline = BatchPDBPreparationPipeline(config_file, output_dir)
    results = pipeline.run_batch_processing()
    status = "completed" if results.get("successful", 0) == results.get("total", 0) else "partial"
    update_artifacts(root, batch_output_dir=str(Path(output_dir).expanduser().resolve()))
    return _result(
        "pdb.batch",
        status,
        root,
        inputs={"config_file": config_file},
        outputs=results,
    )


def run_autodock_prepare(
    target: str,
    receptors_input: Optional[str],
    ligands_input: Optional[str],
    receptors_output: Optional[str],
    ligands_output: Optional[str],
    force_field: str = "AMBER",
    ph: float = 7.4,
    ligand_preparation_backend: str = "engine_aware_full",
    selected_engines: Optional[List[str]] = None,
    autodocktools_prepare_ligand4: Optional[str] = None,
    autodocktools_prepare_receptor4: Optional[str] = None,
    autodocktools_python: Optional[str] = None,
) -> WorkflowStepResult:
    from autodock_preparation import AutoDockPreparationPipeline, PreparationConfig
    from docking.models import (
        DEFAULT_PREPARATION_PH,
        MAX_PREPARATION_PH,
        MIN_PREPARATION_PH,
        validate_ligand_preparation_profile,
        validate_preparation_ph,
    )
    from docking.preparation.ligand_preparation import validate_ligand_preparation_output_contract

    def _count_input_assets(input_path: Optional[str], suffixes: Tuple[str, ...]) -> int:
        if not input_path:
            return 0
        root_path = Path(input_path).expanduser().resolve()
        if not root_path.exists():
            return 0
        if root_path.is_file():
            return int(root_path.suffix.lower() in suffixes)
        count = 0
        for suffix in suffixes:
            count += sum(1 for _ in root_path.rglob(f"*{suffix}"))
        return count

    def _count_prepared_pdbqt(output_path: Optional[str]) -> int:
        if not output_path:
            return 0
        root_path = Path(output_path).expanduser().resolve()
        if not root_path.exists():
            return 0
        return sum(1 for _ in root_path.rglob("*.pdbqt"))

    def _collect_step_reports(ligands_output_dir: str) -> List[Dict[str, object]]:
        reports: List[Dict[str, object]] = []
        report_dir = Path(ligands_output_dir).expanduser().resolve() / "preparation_steps"
        if not report_dir.exists():
            return reports
        for report_file in sorted(report_dir.glob("*.json")):
            try:
                payload = json.loads(report_file.read_text(encoding="utf-8"))
            except Exception as exc:
                reports.append(
                    {
                        "ligand": report_file.stem,
                        "input_file": "",
                        "output_file": "",
                        "report_file": str(report_file),
                        "preparation_method": "",
                        "requested_profile": "",
                        "effective_profile": "",
                        "selected_engines": "",
                        "protonation_ph": None,
                        "is_valid": False,
                        "errors": f"invalid_report_json:{exc}",
                        "warnings": "",
                    }
                )
                continue

            contract = payload.get("output_contract_validation")
            if not isinstance(contract, dict):
                try:
                    contract = validate_ligand_preparation_output_contract(payload)
                except Exception as exc:
                    contract = {
                        "is_valid": False,
                        "errors": [f"contract_validation_error:{exc}"],
                        "warnings": [],
                    }

            input_file = str(payload.get("input_file") or "")
            output_file = str(payload.get("output_file") or "")
            ligand_label = Path(input_file).stem or Path(output_file).stem or report_file.stem
            errors = [str(item) for item in (contract.get("errors") or [])]
            warnings = [str(item) for item in (contract.get("warnings") or [])]
            selected = payload.get("selected_engines") or contract.get("selected_engines") or []
            reports.append(
                {
                    "ligand": ligand_label,
                    "input_file": input_file,
                    "output_file": output_file,
                    "report_file": str(report_file),
                    "preparation_method": str(payload.get("preparation_method") or ""),
                    "requested_profile": str(
                        payload.get("requested_profile")
                        or payload.get("requested_backend")
                        or contract.get("requested_profile")
                        or ""
                    ),
                    "effective_profile": str(payload.get("effective_profile") or contract.get("effective_profile") or ""),
                    "selected_engines": ",".join([str(token).strip().lower() for token in selected if str(token).strip()]),
                    "protonation_ph": payload.get("protonation_ph"),
                    "is_valid": bool(contract.get("is_valid", False)),
                    "errors": "; ".join(errors),
                    "warnings": "; ".join(warnings),
                }
            )
        return reports

    root = _safe_root(output_dir=receptors_output or ligands_output)
    ph_ok, normalized_ph, ph_error = validate_preparation_ph(ph)
    if not ph_ok:
        return _result(
            target,
            "blocked",
            root,
            inputs={
                "receptors_input": receptors_input or "",
                "ligands_input": ligands_input or "",
                "requested_ph": ph,
            },
            outputs={
                "ph_validation": {
                    "is_valid": False,
                    "requested_ph": ph,
                    "normalized_ph": float(DEFAULT_PREPARATION_PH),
                    "min": float(MIN_PREPARATION_PH),
                    "max": float(MAX_PREPARATION_PH),
                    "error": ph_error,
                }
            },
            notes=[ph_error],
        )

    with tempfile.TemporaryDirectory(prefix="pdb_wizard_empty_assets_") as tmp_dir:
        empty_dir = Path(tmp_dir)
        cfg = PreparationConfig(
            ligands_input=str(Path(ligands_input).expanduser().resolve()) if ligands_input else str(empty_dir),
            receptors_input=str(Path(receptors_input).expanduser().resolve()) if receptors_input else str(empty_dir),
            ligands_output=str(Path(ligands_output or empty_dir).expanduser().resolve()),
            receptors_output=str(Path(receptors_output or empty_dir).expanduser().resolve()),
            force_field=force_field,
            ph=normalized_ph,
            ligand_preparation_backend=str(ligand_preparation_backend or "engine_aware_full").strip().lower() or "engine_aware_full",
            ligand_preparation_profile=str(ligand_preparation_backend or "engine_aware_full").strip().lower() or "engine_aware_full",
            selected_engines=[str(engine).strip().lower() for engine in (selected_engines or []) if str(engine).strip()],
            autodocktools_prepare_ligand4=str(autodocktools_prepare_ligand4 or "").strip(),
            autodocktools_prepare_receptor4=str(autodocktools_prepare_receptor4 or "").strip(),
            autodocktools_python=str(autodocktools_python or "").strip(),
        )
        compatibility = validate_ligand_preparation_profile(
            cfg.ligand_preparation_profile,
            cfg.selected_engines,
        )
        if not compatibility.is_valid:
            return _result(
                target,
                "blocked",
                root,
                inputs={"receptors_input": receptors_input or "", "ligands_input": ligands_input or ""},
                outputs={"ligand_preparation_compatibility": compatibility.to_dict()},
                notes=compatibility.errors,
            )
        compatibility_notes = list(compatibility.warnings)
        pipeline = AutoDockPreparationPipeline(cfg)
        deps_ok, missing = pipeline.check_dependencies()
        if not deps_ok:
            return _result(
                target,
                "blocked",
                root,
                inputs={"receptors_input": receptors_input or "", "ligands_input": ligands_input or ""},
                outputs={"missing_dependencies": missing},
                notes=compatibility_notes + ["AutoDock preparation dependencies are missing."],
            )
        success = pipeline.run_enhanced_preparation()
        requested_ligands = target in {"pdb.prepare_both", "pdb.prepare_ligand"}
        requested_receptors = target in {"pdb.prepare_both", "pdb.prepare_protein"}
        ligand_inputs = _count_input_assets(ligands_input, (".sdf", ".mol", ".mol2", ".pdb")) if requested_ligands else 0
        receptor_inputs = _count_input_assets(receptors_input, (".pdb",)) if requested_receptors else 0
        ligand_outputs = _count_prepared_pdbqt(cfg.ligands_output) if requested_ligands else 0
        receptor_outputs = _count_prepared_pdbqt(cfg.receptors_output) if requested_receptors else 0

        notes: List[str] = list(compatibility_notes)
        if requested_ligands and ligand_inputs == 0:
            notes.append(f"No raw ligand files found under {cfg.ligands_input}.")
        if requested_receptors and receptor_inputs == 0:
            notes.append(f"No raw receptor files found under {cfg.receptors_input}.")
        if requested_ligands and ligand_inputs > 0 and ligand_outputs == 0:
            notes.append("Ligand preparation produced 0 PDBQT outputs from non-empty input.")
        if requested_receptors and receptor_inputs > 0 and receptor_outputs == 0:
            notes.append("Receptor preparation produced 0 PDBQT outputs from non-empty input.")

        if not success:
            status = "failed"
        else:
            requested_input_total = ligand_inputs + receptor_inputs
            requested_output_total = ligand_outputs + receptor_outputs
            if requested_input_total == 0:
                status = "blocked"
            elif requested_output_total == 0:
                status = "failed"
            elif notes:
                status = "partial"
            else:
                status = "completed"

        validation_columns = [
            "ligand",
            "input_file",
            "output_file",
            "report_file",
            "preparation_method",
            "requested_profile",
            "effective_profile",
            "selected_engines",
            "protonation_ph",
            "is_valid",
            "errors",
            "warnings",
        ]
        validation_rows = _collect_step_reports(cfg.ligands_output) if requested_ligands else []
        invalid_contract_rows = [row for row in validation_rows if not bool(row.get("is_valid", False))]
        if requested_ligands and ligand_outputs > 0 and not validation_rows:
            notes.append("No per-ligand preparation step reports were found under ligands output/preparation_steps.")
        if invalid_contract_rows:
            notes.append(
                f"Ligand preparation output contract validation flagged {len(invalid_contract_rows)} ligand(s)."
            )
            if status == "completed":
                status = "partial"

        preflight_dir = root / ".workflow" / "preparation"
        preflight_dir.mkdir(parents=True, exist_ok=True)

        validation_df = pd.DataFrame(validation_rows, columns=validation_columns)
        validation_csv = preflight_dir / "ligand_output_contract_validation.csv"
        validation_json = preflight_dir / "ligand_output_contract_validation.json"
        validation_df.to_csv(validation_csv, index=False)
        validation_json.write_text(
            json.dumps(
                {
                    "generated_at_utc": datetime.utcnow().isoformat() + "Z",
                    "requested_target": target,
                    "rows": validation_rows,
                    "invalid_rows": invalid_contract_rows,
                    "counts": {
                        "total_reports": len(validation_rows),
                        "valid_reports": int(validation_df["is_valid"].sum()) if not validation_df.empty else 0,
                        "invalid_reports": len(invalid_contract_rows),
                    },
                },
                indent=2,
            ),
            encoding="utf-8",
        )

        preflight_summary = {
            "generated_at_utc": datetime.utcnow().isoformat() + "Z",
            "target": target,
            "status": status,
            "force_field": str(force_field),
            "requested_ph": float(ph),
            "normalized_ph": float(normalized_ph),
            "ph_min": float(MIN_PREPARATION_PH),
            "ph_max": float(MAX_PREPARATION_PH),
            "ligand_preparation_profile": str(cfg.ligand_preparation_profile),
            "selected_engines": ",".join(cfg.selected_engines),
            "requested_ligands": bool(requested_ligands),
            "requested_receptors": bool(requested_receptors),
            "ligand_input_count": int(ligand_inputs),
            "receptor_input_count": int(receptor_inputs),
            "ligand_prepared_count": int(ligand_outputs),
            "receptor_prepared_count": int(receptor_outputs),
            "step_reports_count": int(len(validation_rows)),
            "step_reports_invalid_count": int(len(invalid_contract_rows)),
            "step_reports_valid_count": int(len(validation_rows) - len(invalid_contract_rows)),
            "notes_count": int(len(notes)),
        }
        preflight_summary_csv = preflight_dir / "preparation_preflight_summary.csv"
        preflight_summary_json = preflight_dir / "preparation_preflight_summary.json"
        pd.DataFrame([preflight_summary]).to_csv(preflight_summary_csv, index=False)
        preflight_summary_json.write_text(
            json.dumps(
                {
                    **preflight_summary,
                    "notes": notes,
                    "ligand_preparation_compatibility": compatibility.to_dict(),
                    "validation_reports_file": str(validation_csv),
                    "validation_reports_json": str(validation_json),
                },
                indent=2,
            ),
            encoding="utf-8",
        )
        outputs = {
            "receptors_output": cfg.receptors_output,
            "ligands_output": cfg.ligands_output,
            "ligand_input_count": ligand_inputs,
            "receptor_input_count": receptor_inputs,
            "ligand_prepared_count": ligand_outputs,
            "receptor_prepared_count": receptor_outputs,
            "ligand_preparation_compatibility": compatibility.to_dict(),
            "ph_validation": {
                "is_valid": True,
                "requested_ph": float(ph),
                "normalized_ph": float(normalized_ph),
                "min": float(MIN_PREPARATION_PH),
                "max": float(MAX_PREPARATION_PH),
                "error": "",
            },
            "ligand_output_contract_validation_csv": str(validation_csv),
            "ligand_output_contract_validation_json": str(validation_json),
            "preparation_preflight_summary_csv": str(preflight_summary_csv),
            "preparation_preflight_summary_json": str(preflight_summary_json),
            "invalid_ligand_output_contracts": len(invalid_contract_rows),
        }
        update_artifacts(
            root,
            prepared_receptors_dir=cfg.receptors_output,
            prepared_ligands_dir=cfg.ligands_output,
            ligand_output_contract_validation_csv=str(validation_csv),
            ligand_output_contract_validation_json=str(validation_json),
            preparation_preflight_summary_csv=str(preflight_summary_csv),
            preparation_preflight_summary_json=str(preflight_summary_json),
        )
        return _result(
            target,
            status,
            root,
            inputs={
                "force_field": force_field,
                "ph": float(normalized_ph),
                "ligand_preparation_backend": cfg.ligand_preparation_profile,
                "selected_engines": cfg.selected_engines,
                "autodocktools_prepare_ligand4": cfg.autodocktools_prepare_ligand4,
                "autodocktools_prepare_receptor4": cfg.autodocktools_prepare_receptor4,
                "autodocktools_python": cfg.autodocktools_python,
            },
            outputs=outputs,
            notes=notes,
        )


def run_prepare_project(argv: List[str]) -> WorkflowStepResult:
    from docking.cli import prepare_docking_main
    from docking.project_layout import manifest_path

    output_dir = _extract_arg_value(argv, "--output")
    root = _safe_root(output_dir=output_dir)
    exit_code = prepare_docking_main(argv)
    status = "completed" if exit_code == 0 else "failed"
    outputs = {"project_root": str(Path(output_dir).expanduser().resolve())} if output_dir else {}
    if output_dir:
        _update_project_context_from_manifest(Path(output_dir).expanduser().resolve())
        update_artifacts(root, project_manifest=str(manifest_path(Path(output_dir).expanduser().resolve())))
    return _result("prep.project", status, root, inputs={"argv": argv}, outputs=outputs)


def run_docking(argv: List[str], engine_target: Optional[str] = None) -> WorkflowStepResult:
    from docking.cli import dock_main

    project_dir = _extract_arg_value(argv, "--project-dir")
    root = _safe_root(project_dir=project_dir)
    exit_code = dock_main(argv)
    status = "completed" if exit_code == 0 else "failed"
    engines = _extract_arg_value(argv, "--engines") or ""
    update_context(root, selected_engines=[e.strip() for e in engines.split(",") if e.strip()])
    if favorite := _extract_arg_value(argv, "--favorite-engine"):
        update_context(root, favorite_engine=favorite)
    if project_dir and is_canonical_project(project_dir):
        _update_project_context_from_manifest(root)
    target = engine_target or "dock.run"
    return _result(target, status, root, inputs={"argv": argv}, outputs={"project_dir": project_dir or ""})


def run_docking_deploy(argv: List[str]) -> WorkflowStepResult:
    from docking.cli import deploy_main
    from docking.project_layout import load_manifest

    project_dir = _extract_arg_value(argv, "--project-dir")
    root = _safe_root(project_dir=project_dir)
    exit_code = deploy_main(argv)
    status = "completed" if exit_code == 0 else "failed"
    outputs = {"project_dir": project_dir or ""}
    if project_dir and is_canonical_project(project_dir):
        manifest = load_manifest(root)
        update_context(root, selected_engines=list(manifest.get("engines", [])))
        update_artifacts(
            root,
            deployment_root=manifest.get("deployment_root", "") or "",
            latest_rerun_manifest_file=manifest.get("latest_rerun_manifest_file", "") or "",
        )
        outputs.update(
            {
                "deployment_root": manifest.get("deployment_root", "") or "",
                "latest_rerun_manifest_file": manifest.get("latest_rerun_manifest_file", "") or "",
            }
        )
    return _result("dock.deploy", status, root, inputs={"argv": argv}, outputs=outputs)


def run_docking_sync(argv: List[str]) -> WorkflowStepResult:
    from docking.cli import sync_main

    project_dir = _extract_arg_value(argv, "--project-dir")
    root = _safe_root(project_dir=project_dir)
    exit_code = sync_main(argv)
    status = "completed" if exit_code == 0 else "failed"
    outputs = {
        "project_dir": project_dir or "",
        "ssh_target": _extract_arg_value(argv, "--ssh-target") or "",
        "remote_project_dir": _extract_arg_value(argv, "--remote-project-dir") or "",
    }
    if outputs["ssh_target"] or outputs["remote_project_dir"]:
        update_artifacts(
            root,
            remote_ssh_target=outputs["ssh_target"],
            remote_project_dir=outputs["remote_project_dir"],
        )
    return _result("dock.sync", status, root, inputs={"argv": argv}, outputs=outputs)


def run_docking_submit(argv: List[str]) -> WorkflowStepResult:
    from docking.cli import submit_main

    project_dir = _extract_arg_value(argv, "--project-dir")
    root = _safe_root(project_dir=project_dir)
    exit_code = submit_main(argv)
    status = "completed" if exit_code == 0 else "failed"
    outputs = {
        "project_dir": project_dir or "",
        "ssh_target": _extract_arg_value(argv, "--ssh-target") or "",
        "remote_project_dir": _extract_arg_value(argv, "--remote-project-dir") or "",
        "round": _extract_arg_value(argv, "--round") or "",
        "mode": _extract_arg_value(argv, "--mode") or "",
    }
    return _result("dock.submit", status, root, inputs={"argv": argv}, outputs=outputs)


def run_analysis_comparative(
    project_dir: str,
    output_dir: Optional[str] = None,
    engines: Optional[List[str]] = None,
    engine_preset: Optional[str] = None,
    promote_exhaustive: bool = False,
    rerun_engine: Optional[str] = None,
    top_per_protein: int = 1,
    winner_only: bool = False,
    min_affinity_advantage: float = 0.0,
    max_rerun_pairs: int = 0,
    pair_allowlist: Optional[str] = None,
    consensus_mode: str = "dockbox_geometric",
    rescoring_scope: str = "top_n_per_protein",
    rescoring_top_n: int = 3,
    prompt_protein_names: bool = False,
    prompt_ligand_names: bool = False,
    complex_query: Optional[str] = None,
    shared_mapping_dir: Optional[str] = None,
    exclude_problematic_ligands: bool = False,
    positive_affinity_threshold: float = 0.0,
    minimum_pose_count: int = 1,
    rmsd_workers: int = 0,
    analysis_scope: str = "full",
    normalization_method: str = "per_engine_rank",
    biology_file: Optional[str] = None,
    biology_mapping_mode: str = "auto",
    hit_class_policy: str = "target_percentile",
    hit_class_strong_percentile: float = 0.10,
    hit_class_moderate_percentile: float = 0.35,
    top_pose_selection_policy: str = "best_affinity",
    top_pose_global_aggregation: str = "best_target",
    best_pose_selection_metric: str = "auto",
    config_file: Optional[str] = None,
) -> WorkflowStepResult:
    from post_docking_analysis.unified_pipeline import UnifiedPostDockingPipeline
    from docking.project_layout import load_manifest, save_manifest, post_docking_root

    root = _safe_root(project_dir=project_dir)
    resolved_config_file, config_overrides, _ = _resolve_analysis_config_overrides(
        project_dir=project_dir,
        config_file=config_file,
    )
    hit_classification = _config_by_path(config_overrides, "quality_control.hit_classification", {})
    if isinstance(hit_classification, dict):
        if hit_class_policy == "target_percentile":
            candidate = str(hit_classification.get("policy") or "").strip().lower()
            if candidate:
                hit_class_policy = candidate
        if float(hit_class_strong_percentile) == 0.10 and "strong_percentile" in hit_classification:
            hit_class_strong_percentile = float(hit_classification["strong_percentile"])
        if float(hit_class_moderate_percentile) == 0.35 and "moderate_percentile" in hit_classification:
            hit_class_moderate_percentile = float(hit_classification["moderate_percentile"])
    output = str(_resolve_analysis_output(Path(project_dir), output_dir, "analyze_comparative"))
    pipeline = UnifiedPostDockingPipeline(
        project_dir=project_dir,
        output_dir=output,
        analysis_mode="comparative_all_engines",
        engines=engines,
        engine_preset=engine_preset,
        favorite_engine=rerun_engine,
        promote_exhaustive=promote_exhaustive,
        rerun_engine=rerun_engine,
        top_per_protein=top_per_protein,
        winner_only=winner_only,
        min_affinity_advantage=min_affinity_advantage,
        max_rerun_pairs=max_rerun_pairs,
        pair_allowlist=pair_allowlist,
        consensus_mode=consensus_mode,
        rescoring_scope=rescoring_scope,
        rescoring_top_n=rescoring_top_n,
        analysis_scope=analysis_scope,
        normalization_method=normalization_method,
        biology_file=biology_file,
        biology_mapping_mode=biology_mapping_mode,
        hit_class_policy=hit_class_policy,
        hit_class_strong_percentile=hit_class_strong_percentile,
        hit_class_moderate_percentile=hit_class_moderate_percentile,
        top_pose_selection_policy=top_pose_selection_policy,
        top_pose_global_aggregation=top_pose_global_aggregation,
        best_pose_selection_metric=best_pose_selection_metric,
        prompt_protein_names=prompt_protein_names,
        prompt_ligand_names=prompt_ligand_names,
        complex_query=complex_query,
        shared_mapping_dir=shared_mapping_dir,
        exclude_problematic_ligands=exclude_problematic_ligands,
        positive_affinity_threshold=positive_affinity_threshold,
        minimum_pose_count=minimum_pose_count,
        rmsd_workers=rmsd_workers,
    )
    success = pipeline.run()
    effective_analysis_mode = str(getattr(pipeline, "analysis_mode", "comparative_all_engines") or "comparative_all_engines")
    effective_single_engine = str(
        getattr(pipeline, "engine", "") or getattr(pipeline, "favorite_engine", "") or ""
    ).strip().lower()
    single_engine_root = Path(output) / effective_single_engine if effective_single_engine else Path(output)
    resolved_top_pose_policy = str(
        getattr(pipeline, "top_pose_selection_policy", top_pose_selection_policy) or "best_affinity"
    )
    resolved_top_pose_aggregation = str(
        getattr(pipeline, "top_pose_global_aggregation", top_pose_global_aggregation) or "best_target"
    )
    manifest = load_manifest(root)
    manifest["top_pose_selection_policy"] = resolved_top_pose_policy
    manifest["top_pose_global_aggregation"] = resolved_top_pose_aggregation
    manifest["best_pose_selection_metric"] = str(
        getattr(pipeline, "best_pose_selection_metric", best_pose_selection_metric) or "auto"
    )
    if pipeline.rerun_manifest_file:
        manifest["latest_rerun_manifest_file"] = pipeline.rerun_manifest_file
    save_manifest(root, manifest)
    update_context(
        root,
        project_root=str(Path(project_dir).expanduser().resolve()),
        analysis_output_dir=output,
        top_pose_selection_policy=resolved_top_pose_policy,
        top_pose_global_aggregation=resolved_top_pose_aggregation,
        best_pose_selection_metric=str(
            getattr(pipeline, "best_pose_selection_metric", best_pose_selection_metric) or "auto"
        ),
    )
    update_artifacts(
        root,
        comparative_output_dir=output,
        comparative_effective_mode=effective_analysis_mode,
        latest_rerun_manifest_file=getattr(pipeline, "rerun_manifest_file", "") or "",
        top_pose_per_ligand_global_file=str(
            getattr(pipeline, "top_pose_outputs", {}).get(
                "top_pose_per_ligand_global_file",
                Path(output) / "top_pose_ligand_performance" / "top_pose_per_ligand_global.csv",
            )
        ),
        top_pose_selection_manifest_file=str(
            getattr(pipeline, "top_pose_outputs", {}).get(
                "top_pose_selection_manifest_file",
                Path(output) / "top_pose_ligand_performance" / "top_pose_selection_manifest.json",
            )
        ),
        top_pose_canonical_root=str(getattr(pipeline, "top_pose_outputs", {}).get("top_pose_canonical_root", "")),
    )
    return _result(
        "analyze.comparative",
        "completed" if success else "failed",
        root,
        inputs={
            "project_dir": project_dir,
            "engines": ",".join(engines or []),
            "engine_preset": engine_preset or "",
            "promote_exhaustive": promote_exhaustive,
            "rerun_engine": rerun_engine or "",
            "top_per_protein": top_per_protein,
            "winner_only": winner_only,
            "min_affinity_advantage": min_affinity_advantage,
            "max_rerun_pairs": max_rerun_pairs,
            "pair_allowlist": pair_allowlist or "",
            "consensus_mode": consensus_mode,
            "rescoring_scope": rescoring_scope,
            "rescoring_top_n": rescoring_top_n,
            "prompt_protein_names": prompt_protein_names,
            "prompt_ligand_names": prompt_ligand_names,
            "complex_query": complex_query or "",
            "exclude_problematic_ligands": exclude_problematic_ligands,
            "positive_affinity_threshold": positive_affinity_threshold,
            "minimum_pose_count": minimum_pose_count,
            "rmsd_workers": rmsd_workers,
            "analysis_scope": analysis_scope,
            "normalization_method": normalization_method,
            "biology_file": biology_file or "",
            "biology_mapping_mode": biology_mapping_mode,
            "hit_class_policy": hit_class_policy,
            "hit_class_strong_percentile": hit_class_strong_percentile,
            "hit_class_moderate_percentile": hit_class_moderate_percentile,
            "top_pose_selection_policy": resolved_top_pose_policy,
            "top_pose_global_aggregation": resolved_top_pose_aggregation,
            "analysis_config_file": resolved_config_file,
        },
        outputs={
            "output_dir": output,
            "effective_analysis_mode": effective_analysis_mode,
            "effective_single_engine": effective_single_engine,
            "engines_in_scope": ",".join(getattr(pipeline, "engines_in_scope", []) or []),
            "rerun_manifest_file": getattr(pipeline, "rerun_manifest_file", "") or "",
            "consensus_mode": consensus_mode,
            "analysis_scope": analysis_scope,
            "normalization_method": normalization_method,
            "rescoring_scope": rescoring_scope,
            "rescoring_top_n": rescoring_top_n,
            "consensus_ranked_hits_file": str(Path(output) / "reports" / "consensus_ranked_hits.csv")
            if effective_analysis_mode != "single_engine"
            else str(single_engine_root / "best_poses.csv"),
            "rescoring_candidates_file": str(Path(output) / "reports" / "rescoring_candidates.csv")
            if effective_analysis_mode != "single_engine"
            else "",
            "consensus_explainability_file": str(Path(output) / "reports" / "consensus_explainability.json")
            if effective_analysis_mode != "single_engine"
            else "",
            "engine_rank_correlation_per_protein_file": str(Path(output) / "reports" / "engine_rank_correlation_per_protein.csv")
            if effective_analysis_mode != "single_engine"
            else "",
            "engine_rank_correlation_global_file": str(Path(output) / "reports" / "engine_rank_correlation_global.csv")
            if effective_analysis_mode != "single_engine"
            else "",
            "canonical_scores_root": str(_canonical_working_scores_root(Path(project_dir))),
            "consensus_ranked_hits_with_classes_file": str(Path(output) / "reports" / "consensus_ranked_hits_with_classes.csv")
            if effective_analysis_mode != "single_engine"
            else "",
            "biology_mapping_report_file": str(Path(output) / "reports" / "biology_mapping_report.json")
            if effective_analysis_mode != "single_engine"
            else "",
            "biology_correlation_global_file": str(Path(output) / "reports" / "biology_correlation_global.csv")
            if effective_analysis_mode != "single_engine"
            else "",
            "biology_correlation_per_protein_file": str(Path(output) / "reports" / "biology_correlation_per_protein.csv")
            if effective_analysis_mode != "single_engine"
            else "",
            "top_pose_per_ligand_per_protein_file": str(
                getattr(pipeline, "top_pose_outputs", {}).get(
                    "top_pose_per_ligand_per_protein_file",
                    Path(output) / "top_pose_ligand_performance" / "top_pose_per_ligand_per_protein.csv",
                )
            ),
            "top_pose_per_ligand_global_file": str(
                getattr(pipeline, "top_pose_outputs", {}).get(
                    "top_pose_per_ligand_global_file",
                    Path(output) / "top_pose_ligand_performance" / "top_pose_per_ligand_global.csv",
                )
            ),
            "top_pose_summary_file": str(
                getattr(pipeline, "top_pose_outputs", {}).get(
                    "ligand_performance_summary_file",
                    Path(output) / "top_pose_ligand_performance" / "ligand_performance_summary.csv",
                )
            ),
            "top_pose_manifest_file": str(
                getattr(pipeline, "top_pose_outputs", {}).get(
                    "top_pose_selection_manifest_file",
                    Path(output) / "top_pose_ligand_performance" / "top_pose_selection_manifest.json",
                )
            ),
            "top_pose_canonical_root": str(
                getattr(pipeline, "top_pose_outputs", {}).get("top_pose_canonical_root", "")
            ),
            "top_pose_selection_policy": resolved_top_pose_policy,
            "top_pose_global_aggregation": resolved_top_pose_aggregation,
        },
    )


def run_analysis_favorite(
    project_dir: str,
    favorite_engine: Optional[str],
    output_dir: Optional[str] = None,
    engines: Optional[List[str]] = None,
    engine_preset: Optional[str] = None,
    enable_poseview: bool = False,
    prompt_protein_names: bool = False,
    prompt_ligand_names: bool = False,
    complex_query: Optional[str] = None,
    rmsd_scopes: Optional[object] = None,
    resume_rmsd: bool = True,
    force_global_rmsd: bool = False,
    global_rmsd_defer_threshold: int = 150,
    shared_mapping_dir: Optional[str] = None,
    exclude_problematic_ligands: bool = False,
    positive_affinity_threshold: float = 0.0,
    minimum_pose_count: int = 1,
    rmsd_workers: int = 0,
    speed_profile: str = "standard",
    analysis_scope: str = "full",
    normalization_method: str = "per_engine_rank",
    biology_file: Optional[str] = None,
    biology_mapping_mode: str = "auto",
    hit_class_policy: str = "target_percentile",
    hit_class_strong_percentile: float = 0.10,
    hit_class_moderate_percentile: float = 0.35,
    top_pose_selection_policy: str = "best_affinity",
    top_pose_global_aggregation: str = "best_target",
    best_pose_selection_metric: str = "auto",
    config_file: Optional[str] = None,
) -> WorkflowStepResult:
    from post_docking_analysis.unified_pipeline import UnifiedPostDockingPipeline

    root = _safe_root(project_dir=project_dir)
    resolved_config_file, config_overrides, _ = _resolve_analysis_config_overrides(
        project_dir=project_dir,
        config_file=config_file,
    )
    hit_classification = _config_by_path(config_overrides, "quality_control.hit_classification", {})
    if isinstance(hit_classification, dict):
        if hit_class_policy == "target_percentile":
            candidate = str(hit_classification.get("policy") or "").strip().lower()
            if candidate:
                hit_class_policy = candidate
        if float(hit_class_strong_percentile) == 0.10 and "strong_percentile" in hit_classification:
            hit_class_strong_percentile = float(hit_classification["strong_percentile"])
        if float(hit_class_moderate_percentile) == 0.35 and "moderate_percentile" in hit_classification:
            hit_class_moderate_percentile = float(hit_classification["moderate_percentile"])
    output = str(_resolve_analysis_output(Path(project_dir), output_dir, "analyze_favorite_engine"))
    pipeline = UnifiedPostDockingPipeline(
        project_dir=project_dir,
        output_dir=output,
        analysis_mode="favorite_engine_continue",
        engines=engines,
        engine_preset=engine_preset,
        favorite_engine=favorite_engine,
        complex_query=complex_query,
        prompt_protein_names=prompt_protein_names,
        prompt_ligand_names=prompt_ligand_names,
        enable_poseview=enable_poseview,
        rmsd_scopes=rmsd_scopes,
        resume_rmsd=resume_rmsd,
        force_global_rmsd=force_global_rmsd,
        global_rmsd_defer_threshold=global_rmsd_defer_threshold,
        shared_mapping_dir=shared_mapping_dir,
        exclude_problematic_ligands=exclude_problematic_ligands,
        positive_affinity_threshold=positive_affinity_threshold,
        minimum_pose_count=minimum_pose_count,
        rmsd_workers=rmsd_workers,
        speed_profile=speed_profile,
        analysis_scope=analysis_scope,
        normalization_method=normalization_method,
        biology_file=biology_file,
        biology_mapping_mode=biology_mapping_mode,
        hit_class_policy=hit_class_policy,
        hit_class_strong_percentile=hit_class_strong_percentile,
        hit_class_moderate_percentile=hit_class_moderate_percentile,
        top_pose_selection_policy=top_pose_selection_policy,
        top_pose_global_aggregation=top_pose_global_aggregation,
        best_pose_selection_metric=best_pose_selection_metric,
    )
    success = pipeline.run()
    resolved_top_pose_policy = str(
        getattr(pipeline, "top_pose_selection_policy", top_pose_selection_policy) or "best_affinity"
    )
    resolved_top_pose_aggregation = str(
        getattr(pipeline, "top_pose_global_aggregation", top_pose_global_aggregation) or "best_target"
    )
    from docking.project_layout import load_manifest, save_manifest

    manifest = load_manifest(root)
    manifest["top_pose_selection_policy"] = resolved_top_pose_policy
    manifest["top_pose_global_aggregation"] = resolved_top_pose_aggregation
    manifest["best_pose_selection_metric"] = str(
        getattr(pipeline, "best_pose_selection_metric", best_pose_selection_metric) or "auto"
    )
    save_manifest(root, manifest)
    update_context(
        root,
        project_root=str(Path(project_dir).expanduser().resolve()),
        analysis_output_dir=output,
        favorite_engine=favorite_engine or "",
        top_pose_selection_policy=resolved_top_pose_policy,
        top_pose_global_aggregation=resolved_top_pose_aggregation,
        best_pose_selection_metric=str(
            getattr(pipeline, "best_pose_selection_metric", best_pose_selection_metric) or "auto"
        ),
    )
    update_artifacts(
        root,
        favorite_analysis_output_dir=output,
        top_pose_per_ligand_global_file=str(
            getattr(pipeline, "top_pose_outputs", {}).get(
                "top_pose_per_ligand_global_file",
                Path(output) / "top_pose_ligand_performance" / "top_pose_per_ligand_global.csv",
            )
        ),
        top_pose_selection_manifest_file=str(
            getattr(pipeline, "top_pose_outputs", {}).get(
                "top_pose_selection_manifest_file",
                Path(output) / "top_pose_ligand_performance" / "top_pose_selection_manifest.json",
            )
        ),
        top_pose_canonical_root=str(getattr(pipeline, "top_pose_outputs", {}).get("top_pose_canonical_root", "")),
    )
    return _result(
        "analyze.favorite_engine",
        "completed" if success else "failed",
        root,
        inputs={
            "project_dir": project_dir,
            "favorite_engine": favorite_engine or "",
            "engines": ",".join(engines or []),
            "engine_preset": engine_preset or "",
            "prompt_protein_names": prompt_protein_names,
            "prompt_ligand_names": prompt_ligand_names,
            "complex_query": complex_query or "",
            "rmsd_scopes": ",".join(rmsd_scopes) if isinstance(rmsd_scopes, (list, tuple, set)) else str(rmsd_scopes or ""),
            "resume_rmsd": resume_rmsd,
            "force_global_rmsd": force_global_rmsd,
            "global_rmsd_defer_threshold": global_rmsd_defer_threshold,
            "exclude_problematic_ligands": exclude_problematic_ligands,
            "positive_affinity_threshold": positive_affinity_threshold,
            "minimum_pose_count": minimum_pose_count,
            "rmsd_workers": rmsd_workers,
            "speed_profile": speed_profile,
            "analysis_scope": analysis_scope,
            "normalization_method": normalization_method,
            "biology_file": biology_file or "",
            "biology_mapping_mode": biology_mapping_mode,
            "hit_class_policy": hit_class_policy,
            "hit_class_strong_percentile": hit_class_strong_percentile,
            "hit_class_moderate_percentile": hit_class_moderate_percentile,
            "top_pose_selection_policy": resolved_top_pose_policy,
            "top_pose_global_aggregation": resolved_top_pose_aggregation,
            "analysis_config_file": resolved_config_file,
        },
        outputs={
            "output_dir": output,
            "engines_in_scope": ",".join(getattr(pipeline, "engines_in_scope", []) or []),
            "top_pose_per_ligand_per_protein_file": str(
                getattr(pipeline, "top_pose_outputs", {}).get(
                    "top_pose_per_ligand_per_protein_file",
                    Path(output) / "top_pose_ligand_performance" / "top_pose_per_ligand_per_protein.csv",
                )
            ),
            "top_pose_per_ligand_global_file": str(
                getattr(pipeline, "top_pose_outputs", {}).get(
                    "top_pose_per_ligand_global_file",
                    Path(output) / "top_pose_ligand_performance" / "top_pose_per_ligand_global.csv",
                )
            ),
            "top_pose_summary_file": str(
                getattr(pipeline, "top_pose_outputs", {}).get(
                    "ligand_performance_summary_file",
                    Path(output) / "top_pose_ligand_performance" / "ligand_performance_summary.csv",
                )
            ),
            "top_pose_manifest_file": str(
                getattr(pipeline, "top_pose_outputs", {}).get(
                    "top_pose_selection_manifest_file",
                    Path(output) / "top_pose_ligand_performance" / "top_pose_selection_manifest.json",
                )
            ),
            "top_pose_canonical_root": str(
                getattr(pipeline, "top_pose_outputs", {}).get("top_pose_canonical_root", "")
            ),
            "top_pose_selection_policy": resolved_top_pose_policy,
            "top_pose_global_aggregation": resolved_top_pose_aggregation,
        },
    )


def _shell_command_line(tokens: List[str]) -> str:
    return " ".join(shlex.quote(str(token)) for token in tokens)


def _latest_bundle_root(dataset_root: Path) -> Optional[Path]:
    bundles = [path for path in dataset_root.glob("gnina_selected_pose_pdb_export_*") if path.is_dir()]
    if not bundles:
        return None
    return max(bundles, key=lambda path: path.stat().st_mtime)


def _latest_named_artifact(root: Path, filename: str) -> str:
    candidates: List[Path] = []
    try:
        candidates = [path for path in root.rglob(filename) if path.is_file()]
    except Exception:
        return ""
    if not candidates:
        return ""
    latest = max(candidates, key=lambda path: path.stat().st_mtime)
    return str(latest)


def _run_clean_interaction_pipeline_target(
    *,
    project_dir: Optional[str],
    output_dir: Optional[str],
    dataset_root: Optional[str],
    python_bin: Optional[str],
    pipeline_execution_mode: str,
    target_filter: Optional[str],
    ligand_filter: Optional[str],
    max_targets: int,
    max_ligands: int,
    max_poses: int,
    run_chemistry_fixes: bool,
    run_layered_plip: bool,
    run_plip: bool,
    run_prolif: bool,
    dry_run: bool,
) -> WorkflowStepResult:
    target = "analyze.interactions.clean"
    root = _safe_root(project_dir=project_dir)
    dataset = Path(dataset_root or DEFAULT_CLEAN_INTERACTION_DATASET_ROOT).expanduser().resolve()
    run_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_root = (
        Path(output_dir).expanduser().resolve()
        if output_dir
        else dataset / "clean_interaction_pipeline_runs" / run_id
    )
    try:
        output_root.mkdir(parents=True, exist_ok=True)
    except Exception as exc:
        notes = [
            f"Could not create clean-interaction output directory: {output_root}",
            f"error={exc}",
            "Provide --output to a writable path and retry.",
        ]
        return _result(
            target,
            "failed",
            root,
            inputs={
                "project_dir": project_dir or "",
                "dataset_root": str(dataset),
                "output_dir": str(output_root),
            },
            outputs={"output_dir": str(output_root)},
            notes=notes,
        )

    notes: List[str] = []
    if not dataset.exists():
        notes.append(f"Dataset root not found: {dataset}")
        return _result(
            target,
            "failed",
            root,
            inputs={"dataset_root": str(dataset)},
            outputs={"output_dir": str(output_root)},
            notes=notes,
        )

    py_exec = str(python_bin or "python")
    mode = str(pipeline_execution_mode or "local").strip() or "local"
    if mode not in {"local", "local_dry_run", "hpc"}:
        mode = "local"
        notes.append("Invalid execution mode provided; defaulted to local")

    step_defs: List[Dict[str, object]] = []

    def _append_step(label: str, cmd: List[str], cwd: Path) -> None:
        step_defs.append(
            {
                "label": label,
                "cmd": [str(token) for token in cmd],
                "cwd": str(cwd),
                "command_line": _shell_command_line([str(token) for token in cmd]),
            }
        )

    if run_chemistry_fixes:
        fix_exports_script = dataset / "fix_gnina_complex_exports.py"
        if fix_exports_script.exists():
            _append_step("fix_gnina_complex_exports", [py_exec, str(fix_exports_script)], dataset)
        else:
            notes.append(f"Chemistry-fix script missing: {fix_exports_script}")

        latest_bundle = _latest_bundle_root(dataset)
        clean_script = dataset / "6-Manuscript_" / "closing_thesis" / "scripts" / "clean_pdb_complexes_for_plip.py"
        if latest_bundle is None:
            notes.append("No gnina_selected_pose_pdb_export_* bundle found; skipped clean_pdb_complexes_for_plip stage")
        elif not clean_script.exists():
            notes.append(f"Chemistry-fix script missing: {clean_script}")
        else:
            clean_source = latest_bundle / "complexes"
            clean_output = latest_bundle / "complexes_plip_cleaned"
            if clean_source.exists():
                _append_step(
                    "clean_pdb_complexes_for_plip",
                    [
                        py_exec,
                        str(clean_script),
                        "--source",
                        str(clean_source),
                        "--output",
                        str(clean_output),
                        "--note",
                        "Cleaned for PLIP via OmniDock analyze interactions clean",
                    ],
                    dataset,
                )
            else:
                notes.append(f"Expected complexes directory missing for cleaning stage: {clean_source}")

    if run_layered_plip:
        layered_script = dataset / "extract_layered_plip_interactions.py"
        if layered_script.exists():
            _append_step("extract_layered_plip_interactions", [py_exec, str(layered_script), "--run-plip"], dataset)
        else:
            notes.append(f"Layered PLIP script missing: {layered_script}")

    if run_plip or run_prolif:
        pose_cmd = [
            py_exec,
            "-m",
            "pose_ensemble_appraisal",
            "run",
            "--project-root",
            str(dataset),
            "--dataset-root",
            str(dataset),
            "--execution-mode",
            mode,
        ]
        if run_plip:
            pose_cmd.append("--run-plip")
        if run_prolif:
            pose_cmd.append("--run-prolif")
        if target_filter:
            pose_cmd.extend(["--target-filter", str(target_filter)])
        if ligand_filter:
            pose_cmd.extend(["--ligand-filter", str(ligand_filter)])
        if int(max_targets or 0) > 0:
            pose_cmd.extend(["--max-targets", str(int(max_targets))])
        if int(max_ligands or 0) > 0:
            pose_cmd.extend(["--max-ligands", str(int(max_ligands))])
        if int(max_poses or 0) > 0:
            pose_cmd.extend(["--max-poses", str(int(max_poses))])
        _append_step("pose_ensemble_appraisal_run", pose_cmd, dataset)
    else:
        notes.append("PLIP and ProLIF stages were both disabled; skipped pose_ensemble_appraisal run stage")

    if not step_defs:
        notes.append("No executable clean-interaction steps were assembled")
        return _result(
            target,
            "failed",
            root,
            inputs={
                "project_dir": project_dir or "",
                "dataset_root": str(dataset),
            },
            outputs={"output_dir": str(output_root)},
            notes=notes,
        )

    command_log_file = output_root / "command_log.txt"
    command_manifest_file = output_root / "command_manifest.json"
    command_table_file = output_root / "command_manifest.csv"
    command_set_file = output_root / "run_clean_interaction_pipeline.sh"

    shell_lines = ["#!/usr/bin/env bash", "set -euo pipefail", ""]
    for step in step_defs:
        shell_lines.append(f"( cd {shlex.quote(str(step['cwd']))} && {str(step['command_line'])} )")
    command_set_file.write_text("\n".join(shell_lines) + "\n", encoding="utf-8")
    try:
        command_set_file.chmod(0o755)
    except OSError:
        pass

    command_records: List[Dict[str, object]] = []
    failed_step = ""
    with command_log_file.open("w", encoding="utf-8") as log_handle:
        log_handle.write("# OmniDock clean interaction pipeline command log\n")
        log_handle.write(f"# dataset_root={dataset}\n")
        log_handle.write(f"# dry_run={dry_run}\n\n")

        for index, step in enumerate(step_defs, start=1):
            label = str(step["label"])
            command_line = str(step["command_line"])
            cwd = str(step["cwd"])
            cmd = [str(token) for token in step["cmd"]]  # type: ignore[index]
            started_at = datetime.now().isoformat(timespec="seconds")

            print(f"[{index}/{len(step_defs)}] {label}")
            print(f"    $ {command_line}")

            record: Dict[str, object] = {
                "index": index,
                "label": label,
                "cwd": cwd,
                "command_line": command_line,
                "started_at": started_at,
                "dry_run": bool(dry_run),
                "return_code": 0,
                "status": "planned" if dry_run else "ok",
                "log_file": str(command_log_file),
            }

            log_handle.write(f"## [{index}/{len(step_defs)}] {label}\n")
            log_handle.write(f"$ (cd {cwd} && {command_line})\n")
            log_handle.write(f"started_at={started_at}\n")

            if not dry_run:
                completed = subprocess.run(
                    cmd,
                    cwd=cwd,
                    capture_output=True,
                    text=True,
                )
                record["return_code"] = int(completed.returncode)
                record["status"] = "ok" if completed.returncode == 0 else "failed"
                log_handle.write(completed.stdout or "")
                if completed.stderr:
                    log_handle.write("\n[stderr]\n")
                    log_handle.write(completed.stderr)
                if completed.returncode != 0:
                    failed_step = label
                    command_records.append(record)
                    log_handle.write(f"\nfailed_at={datetime.now().isoformat(timespec='seconds')}\n\n")
                    break
            log_handle.write(f"\nfinished_at={datetime.now().isoformat(timespec='seconds')}\n\n")
            command_records.append(record)

    manifest_payload = {
        "target": target,
        "dataset_root": str(dataset),
        "output_root": str(output_root),
        "dry_run": bool(dry_run),
        "execution_mode": mode,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "steps": step_defs,
        "records": command_records,
        "notes": notes,
        "command_set_script": str(command_set_file),
    }
    command_manifest_file.write_text(json.dumps(manifest_payload, indent=2), encoding="utf-8")
    pd.DataFrame(command_records).to_csv(command_table_file, index=False)

    layered_summary = _latest_named_artifact(dataset, "layered_plip_summary.csv")
    plip_job_manifest = _latest_named_artifact(dataset, "plip_job_manifest.csv")

    status = "completed" if not failed_step else "failed"
    if dry_run:
        notes.append("Dry-run mode: commands were assembled and logged without execution")
    if failed_step:
        notes.append(f"Pipeline aborted at step: {failed_step}")

    update_context(root, analysis_output_dir=str(output_root))
    update_artifacts(
        root,
        last_analysis_output=str(output_root),
        clean_interaction_command_manifest=str(command_manifest_file),
        clean_interaction_command_table=str(command_table_file),
    )

    return _result(
        target,
        status,
        root,
        inputs={
            "project_dir": project_dir or "",
            "dataset_root": str(dataset),
            "python_bin": py_exec,
            "execution_mode": mode,
            "target_filter": target_filter or "",
            "ligand_filter": ligand_filter or "",
            "max_targets": int(max_targets or 0),
            "max_ligands": int(max_ligands or 0),
            "max_poses": int(max_poses or 0),
            "run_chemistry_fixes": bool(run_chemistry_fixes),
            "run_layered_plip": bool(run_layered_plip),
            "run_plip": bool(run_plip),
            "run_prolif": bool(run_prolif),
            "dry_run": bool(dry_run),
        },
        outputs={
            "dataset_root": str(dataset),
            "output_dir": str(output_root),
            "command_log_file": str(command_log_file),
            "command_manifest_file": str(command_manifest_file),
            "command_table_file": str(command_table_file),
            "command_set_script": str(command_set_file),
            "planned_steps": len(step_defs),
            "completed_steps": len(command_records),
            "failed_step": failed_step,
            "layered_plip_summary_file": layered_summary,
            "plip_job_manifest_file": plip_job_manifest,
        },
        notes=notes,
    )


def run_analysis_target(
    target: str,
    project_dir: Optional[str] = None,
    output_dir: Optional[str] = None,
    engine: Optional[str] = None,
    favorite_engine: Optional[str] = None,
    engines: Optional[List[str]] = None,
    engine_preset: Optional[str] = None,
    sdf_folder: Optional[str] = None,
    log_folder: Optional[str] = None,
    receptors_folder: Optional[str] = None,
    pairlist: Optional[str] = None,
    ligplus_root: Optional[str] = None,
    enable_poseview: bool = False,
    prompt_protein_names: bool = False,
    prompt_ligand_names: bool = False,
    complex_query: Optional[str] = None,
    rmsd_scopes: Optional[object] = None,
    resume_rmsd: bool = True,
    force_global_rmsd: bool = False,
    global_rmsd_defer_threshold: int = 150,
    shared_mapping_dir: Optional[str] = None,
    exclude_problematic_ligands: bool = False,
    positive_affinity_threshold: float = 0.0,
    minimum_pose_count: int = 1,
    rmsd_workers: int = 0,
    speed_profile: str = "standard",
    dataset_root: Optional[str] = None,
    python_bin: Optional[str] = None,
    pipeline_execution_mode: str = "local",
    target_filter: Optional[str] = None,
    ligand_filter: Optional[str] = None,
    max_targets: int = 0,
    max_ligands: int = 0,
    max_poses: int = 0,
    run_chemistry_fixes: bool = True,
    run_layered_plip: bool = True,
    run_plip: bool = True,
    run_prolif: bool = True,
    dry_run: bool = False,
    config_file: Optional[str] = None,
) -> WorkflowStepResult:
    if target in {"analyze.comparative", "analyze.favorite_engine"}:
        raise ValueError(f"Target {target} should use the dedicated high-level analysis handler")

    # Contract enforcement: standalone interaction aliases must route through the
    # clean interaction pipeline entrypoint instead of bypassing mandatory stages.
    if target in _INTERACTION_TARGETS_ROUTED_TO_CLEAN:
        if not dataset_root and project_dir:
            dataset_root = project_dir
        target = "analyze.interactions.clean"

    if target in _BLOCKED_LEGACY_STAGE_TARGETS:
        _raise_legacy_target_migration_error(target)

    if target == "analyze.interactions.clean":
        resolved_config_file, config_overrides, _ = _resolve_analysis_config_overrides(
            project_dir=project_dir,
            config_file=config_file,
        )
        clean_cfg = _config_by_path(config_overrides, "clean_interactions", {})
        if isinstance(clean_cfg, dict):
            if int(max_targets or 0) == 0 and clean_cfg.get("max_targets") is not None:
                max_targets = int(clean_cfg.get("max_targets") or 0)
            if int(max_ligands or 0) == 0 and clean_cfg.get("max_ligands") is not None:
                max_ligands = int(clean_cfg.get("max_ligands") or 0)
            if int(max_poses or 0) == 0 and clean_cfg.get("max_poses") is not None:
                max_poses = int(clean_cfg.get("max_poses") or 0)
            if run_chemistry_fixes and clean_cfg.get("run_chemistry_fixes") is not None:
                run_chemistry_fixes = bool(clean_cfg.get("run_chemistry_fixes"))
            if run_layered_plip and clean_cfg.get("run_layered_plip") is not None:
                run_layered_plip = bool(clean_cfg.get("run_layered_plip"))
            if run_plip and clean_cfg.get("run_plip") is not None:
                run_plip = bool(clean_cfg.get("run_plip"))
            if run_prolif and clean_cfg.get("run_prolif") is not None:
                run_prolif = bool(clean_cfg.get("run_prolif"))
        return _run_clean_interaction_pipeline_target(
            project_dir=project_dir,
            output_dir=output_dir,
            dataset_root=dataset_root,
            python_bin=python_bin,
            pipeline_execution_mode=pipeline_execution_mode,
            target_filter=target_filter,
            ligand_filter=ligand_filter,
            max_targets=max_targets,
            max_ligands=max_ligands,
            max_poses=max_poses,
            run_chemistry_fixes=run_chemistry_fixes,
            run_layered_plip=run_layered_plip,
            run_plip=run_plip,
            run_prolif=run_prolif,
            dry_run=dry_run,
        )

    if project_dir and is_canonical_project(project_dir):
        target_engine = engine or favorite_engine or _favorite_engine_from_manifest(project_dir)
        if not target_engine:
            raise ValueError("Canonical multi-engine stage analysis requires --engine or a favorite_engine in project_manifest.json")
        requested_poseview = bool(enable_poseview or target == "analyze.interactions.poseview")
        scope = _analysis_scope_for_stage_target(target)
        delegated = run_analysis_favorite(
            project_dir=project_dir,
            favorite_engine=target_engine,
            engines=engines,
            engine_preset=engine_preset,
            output_dir=output_dir,
            enable_poseview=requested_poseview,
            prompt_protein_names=prompt_protein_names,
            prompt_ligand_names=prompt_ligand_names,
            complex_query=complex_query,
            rmsd_scopes=rmsd_scopes,
            resume_rmsd=resume_rmsd,
            force_global_rmsd=force_global_rmsd,
            global_rmsd_defer_threshold=global_rmsd_defer_threshold,
            shared_mapping_dir=shared_mapping_dir,
            exclude_problematic_ligands=exclude_problematic_ligands,
            positive_affinity_threshold=positive_affinity_threshold,
            minimum_pose_count=minimum_pose_count,
            rmsd_workers=rmsd_workers,
            speed_profile=speed_profile,
            analysis_scope=scope,
            config_file=config_file,
        )
        delegated_inputs = delegated.inputs if isinstance(delegated.inputs, dict) else {}
        root = _safe_root(project_dir=project_dir, output_dir=output_dir)
        notes = [
            "Canonical stage target delegated to unified favorite-engine execution path.",
            f"requested_target={target}",
            f"mapped_analysis_scope={scope}",
            f"engine={target_engine}",
            f"poseview_enabled={requested_poseview}",
        ]
        return _result(
            target,
            str(delegated.status),
            root,
            inputs={
                "project_dir": project_dir,
                "engine": target_engine,
                "analysis_scope": scope,
                "complex_query": complex_query or "",
                "engines": ",".join(engines or []),
                "engine_preset": engine_preset or "",
                "rmsd_scopes": ",".join(rmsd_scopes) if isinstance(rmsd_scopes, (list, tuple, set)) else str(rmsd_scopes or ""),
                "resume_rmsd": resume_rmsd,
                "force_global_rmsd": force_global_rmsd,
                "global_rmsd_defer_threshold": global_rmsd_defer_threshold,
                "exclude_problematic_ligands": exclude_problematic_ligands,
                "positive_affinity_threshold": positive_affinity_threshold,
                "minimum_pose_count": minimum_pose_count,
                "rmsd_workers": rmsd_workers,
                "speed_profile": speed_profile,
                "analysis_config_file": str(delegated_inputs.get("analysis_config_file", "")),
            },
            outputs=dict(delegated.outputs),
            notes=notes,
        )

    raise ValueError(
        "Workflow stage targets now require a canonical DockForge project and unified execution path. "
        "Initialize/upgrade the project with `python main.py workflow init` (or interactive workflow setup), "
        "then rerun the analysis stage. For legacy raw GNINA folders, use `python -m post_docking_analysis.simplified_cli`."
    )


def _resolve_gnina_inputs(
    project_dir: Optional[str],
    sdf_folder: Optional[str],
    log_folder: Optional[str],
    receptors_folder: Optional[str],
    pairlist: Optional[str],
) -> Tuple[Path, Path, Path, Optional[Path]]:
    from post_docking_analysis.gnina_hpc_adapter import detect_gnina_layout

    if project_dir:
        layout = detect_gnina_layout(project_dir)
        resolved_sdf = layout["sdf_folder"]
        resolved_log = layout["log_folder"]
        resolved_receptors = layout["receptors_folder"]
        resolved_pairlist = Path(pairlist).expanduser().resolve() if pairlist else layout["pairlist_file"]
        if not resolved_sdf or not resolved_log or not resolved_receptors:
            raise ValueError("Could not resolve GNINA SDF/log/receptor folders from the provided project directory")
        return resolved_sdf, resolved_log, resolved_receptors, resolved_pairlist

    if not (sdf_folder and log_folder and receptors_folder):
        raise ValueError("GNINA stage analysis requires --project-dir or explicit --sdf-folder, --log-folder, and --receptors-folder")
    return (
        Path(sdf_folder).expanduser().resolve(),
        Path(log_folder).expanduser().resolve(),
        Path(receptors_folder).expanduser().resolve(),
        Path(pairlist).expanduser().resolve() if pairlist else None,
    )


def _build_simplified_pipeline(
    project_dir: Optional[str],
    output_dir: str,
    sdf_folder: Optional[str] = None,
    log_folder: Optional[str] = None,
    receptors_folder: Optional[str] = None,
    pairlist: Optional[str] = None,
    ligplus_root: Optional[str] = None,
    enable_poseview: bool = False,
    prompt_protein_names: bool = False,
    prompt_ligand_names: bool = False,
    complex_query: Optional[str] = None,
    rmsd_scopes: Optional[object] = None,
    resume_rmsd: bool = True,
    force_global_rmsd: bool = False,
    global_rmsd_defer_threshold: int = 150,
    shared_mapping_dir: Optional[str] = None,
    exclude_problematic_ligands: bool = False,
    positive_affinity_threshold: float = 0.0,
    minimum_pose_count: int = 1,
    rmsd_workers: int = 0,
    analysis_config: Optional[Dict[str, object]] = None,
) -> SimplifiedPostDockingPipeline:
    from post_docking_analysis.simplified_pipeline import SimplifiedPostDockingPipeline

    resolved_sdf, resolved_log, resolved_receptors, resolved_pairlist = _resolve_gnina_inputs(
        project_dir,
        sdf_folder,
        log_folder,
        receptors_folder,
        pairlist,
    )
    return SimplifiedPostDockingPipeline(
        sdf_folder=str(resolved_sdf),
        log_folder=str(resolved_log),
        receptors_folder=str(resolved_receptors),
        output_dir=output_dir,
        pairlist_file=str(resolved_pairlist) if resolved_pairlist else None,
        run_rmsd=True,
        run_visualizations=True,
        ligplus_root=ligplus_root,
        prompt_protein_names=prompt_protein_names,
        prompt_ligand_names=prompt_ligand_names,
        enable_poseview=enable_poseview,
        complex_query=complex_query,
        rmsd_scopes=rmsd_scopes,
        resume_rmsd=resume_rmsd,
        force_global_rmsd=force_global_rmsd,
        global_rmsd_defer_threshold=global_rmsd_defer_threshold,
        shared_mapping_dir=shared_mapping_dir,
        exclude_problematic_ligands=exclude_problematic_ligands,
        positive_affinity_threshold=positive_affinity_threshold,
        minimum_pose_count=minimum_pose_count,
        rmsd_workers=rmsd_workers,
        analysis_config=analysis_config,
    )


def _prepare_simplified_context(
    pipeline: SimplifiedPostDockingPipeline,
    need_affinity: bool = False,
    need_polypharmacology: bool = False,
) -> bool:
    steps: List[Tuple[str, object]] = [
        ("Find input files", pipeline._find_input_files),
        ("Generate all_scores.csv", pipeline._generate_scores_csv),
        ("Match poses to receptors", pipeline._match_poses_to_receptors),
        ("Create complex PDB files", pipeline._create_complexes),
    ]
    if need_affinity:
        steps.append(("Binding affinity analysis", pipeline._analyze_binding_affinity))
    if need_polypharmacology:
        steps.append(("Polypharmacology analysis", pipeline._analyze_polypharmacology))

    total_steps = len(steps)
    for index, (label, callback) in enumerate(steps, start=1):
        if getattr(pipeline, "logger", None):
            pipeline.logger.info(f"📍 Context step [{index}/{total_steps}] {label}")
        if not bool(callback()):
            if getattr(pipeline, "logger", None):
                pipeline.logger.error(f"❌ Context step failed [{index}/{total_steps}] {label}")
            return False
    return True


def _run_gnina_analysis_target(
    target: str,
    project_dir: Optional[str] = None,
    output_dir: Optional[str] = None,
    sdf_folder: Optional[str] = None,
    log_folder: Optional[str] = None,
    receptors_folder: Optional[str] = None,
    pairlist: Optional[str] = None,
    ligplus_root: Optional[str] = None,
    enable_poseview: bool = False,
    prompt_protein_names: bool = False,
    prompt_ligand_names: bool = False,
    complex_query: Optional[str] = None,
    rmsd_scopes: Optional[object] = None,
    resume_rmsd: bool = True,
    force_global_rmsd: bool = False,
    global_rmsd_defer_threshold: int = 150,
    shared_mapping_dir: Optional[str] = None,
    exclude_problematic_ligands: bool = False,
    positive_affinity_threshold: float = 0.0,
    minimum_pose_count: int = 1,
    rmsd_workers: int = 0,
    analysis_config: Optional[Dict[str, object]] = None,
    analysis_config_file: str = "",
) -> WorkflowStepResult:
    if target in _BLOCKED_LEGACY_STAGE_TARGETS:
        _raise_legacy_target_migration_error(target)
    if target in _INTERACTION_TARGETS_ROUTED_TO_CLEAN:
        raise ValueError(
            f"Target {target} is an interaction alias. Dispatch it through run_analysis_target "
            "so it is normalized to analyze.interactions.clean and enforces the mandatory clean contract."
        )

    poseview_enabled = bool(enable_poseview or target == "analyze.interactions.poseview")
    root = _safe_root(project_dir=project_dir, output_dir=output_dir)
    out_dir = _resolve_analysis_output(root, output_dir, target)
    pipeline = _build_simplified_pipeline(
        project_dir=project_dir,
        output_dir=str(out_dir),
        sdf_folder=sdf_folder,
        log_folder=log_folder,
        receptors_folder=receptors_folder,
        pairlist=pairlist,
        ligplus_root=ligplus_root,
        enable_poseview=poseview_enabled,
        prompt_protein_names=prompt_protein_names,
        prompt_ligand_names=prompt_ligand_names,
        complex_query=complex_query,
        rmsd_scopes=rmsd_scopes,
        resume_rmsd=resume_rmsd,
        force_global_rmsd=force_global_rmsd,
        global_rmsd_defer_threshold=global_rmsd_defer_threshold,
        shared_mapping_dir=shared_mapping_dir,
        exclude_problematic_ligands=exclude_problematic_ligands,
        positive_affinity_threshold=positive_affinity_threshold,
        minimum_pose_count=minimum_pose_count,
        rmsd_workers=rmsd_workers,
        analysis_config=analysis_config,
    )

    if target == "analyze.stage.hierarchical":
        success = _prepare_simplified_context(pipeline) and pipeline._analyze_binding_affinity()
    elif target == "analyze.stage.polypharmacology":
        success = _prepare_simplified_context(pipeline, need_affinity=True) and pipeline._analyze_polypharmacology()
    elif target == "analyze.stage.rmsd":
        success = _prepare_simplified_context(pipeline, need_affinity=True) and pipeline._analyze_rmsd()
    elif target == "analyze.stage.reports":
        success = _prepare_simplified_context(pipeline) and pipeline._generate_reports()
    elif target == "analyze.stage.visualizations":
        success = _prepare_simplified_context(pipeline, need_affinity=True) and pipeline._generate_visualizations()
    elif target == "analyze.interactions.poseview":
        success = _prepare_simplified_context(pipeline, need_affinity=True) and pipeline._generate_poseview_diagrams()
    elif target == "analyze.visuals.py3dmol":
        success = _prepare_simplified_context(pipeline, need_affinity=True) and pipeline._generate_py3dmol_visualizations()
    else:
        raise ValueError(f"Unsupported GNINA analysis target: {target}")

    if success and hasattr(pipeline, "_organize_output_artifacts"):
        try:
            pipeline._organize_output_artifacts()
        except Exception:
            pass

    _prune_empty_dirs(out_dir)
    update_context(root, analysis_output_dir=str(out_dir))
    update_artifacts(root, last_analysis_output=str(out_dir))
    return _result(
        target,
        "completed" if success else "failed",
        root,
        inputs={
            "project_dir": project_dir or "",
            "sdf_folder": sdf_folder or "",
            "log_folder": log_folder or "",
            "receptors_folder": receptors_folder or "",
            "complex_query": complex_query or "",
            "rmsd_scopes": ",".join(rmsd_scopes) if isinstance(rmsd_scopes, (list, tuple, set)) else str(rmsd_scopes or ""),
            "resume_rmsd": resume_rmsd,
            "force_global_rmsd": force_global_rmsd,
            "global_rmsd_defer_threshold": global_rmsd_defer_threshold,
            "exclude_problematic_ligands": exclude_problematic_ligands,
            "positive_affinity_threshold": positive_affinity_threshold,
            "minimum_pose_count": minimum_pose_count,
            "rmsd_workers": rmsd_workers,
            "analysis_config_file": analysis_config_file,
        },
        outputs={"output_dir": str(out_dir)},
    )


def _prepare_non_gnina_context(
    project_dir: str,
    engine: str,
    output_dir: Path,
    ligplus_root: Optional[str],
    enable_poseview: bool,
    prompt_protein_names: bool,
    prompt_ligand_names: bool,
    complex_query: Optional[str],
    rmsd_scopes: Optional[object],
    resume_rmsd: bool,
    force_global_rmsd: bool,
    global_rmsd_defer_threshold: int,
    shared_mapping_dir: Optional[str],
    exclude_problematic_ligands: bool,
    positive_affinity_threshold: float,
    minimum_pose_count: int,
    rmsd_workers: int,
    speed_profile: str,
    analysis_config: Optional[Dict[str, object]],
) -> Tuple["UnifiedPostDockingPipeline", "SimplifiedPostDockingPipeline", pd.DataFrame, pd.DataFrame, int]:
    from docking.project_layout import ensure_engine_layout, pairlist_path, shared_receptors_dir
    from post_docking_analysis.unified_pipeline import UnifiedPostDockingPipeline
    from post_docking_analysis.protein_naming import build_protein_name_mapping
    from post_docking_analysis.simplified_input_handler import find_receptor_files
    from post_docking_analysis.simplified_pipeline import SimplifiedPostDockingPipeline

    analysis = UnifiedPostDockingPipeline(
        project_dir=project_dir,
        output_dir=str(output_dir),
        analysis_mode="single_engine",
        engine=engine,
        favorite_engine=engine,
        complex_query=complex_query,
        rmsd_scopes=rmsd_scopes,
        resume_rmsd=resume_rmsd,
        force_global_rmsd=force_global_rmsd,
        global_rmsd_defer_threshold=global_rmsd_defer_threshold,
        shared_mapping_dir=shared_mapping_dir,
        exclude_problematic_ligands=exclude_problematic_ligands,
        positive_affinity_threshold=positive_affinity_threshold,
        minimum_pose_count=minimum_pose_count,
        rmsd_workers=rmsd_workers,
        speed_profile=speed_profile,
    )
    scores = analysis._load_or_build_scores()
    engine_scores = scores[scores["engine"] == engine].copy()
    if engine_scores.empty:
        raise ValueError(f"No normalized scores were found for engine '{engine}'")

    downstream_results = analysis._build_downstream_results(engine_scores)
    best_poses = downstream_results["best_poses"].copy()
    raw_dir = output_dir / "raw_data"
    raw_dir.mkdir(parents=True, exist_ok=True)

    # The constants are module-level, not class-level.
    from post_docking_analysis.unified_pipeline import UNIFIED_COMPAT_COLUMNS

    downstream_results["full_data"][UNIFIED_COMPAT_COLUMNS].to_csv(raw_dir / "unified_all_scores.csv", index=False)
    best_poses[UNIFIED_COMPAT_COLUMNS].to_csv(raw_dir / "unified_best_poses.csv", index=False)

    bridge_scores = analysis._build_simplified_bridge_scores(engine_scores)
    bridge_scores.to_csv(output_dir / "all_scores.csv", index=False)

    simplified = SimplifiedPostDockingPipeline(
        sdf_folder=str(ensure_engine_layout(project_dir, engine)["poses"]),
        log_folder=str(ensure_engine_layout(project_dir, engine)["logs"]),
        receptors_folder=str(shared_receptors_dir(project_dir)),
        output_dir=str(output_dir),
        pairlist_file=str(pairlist_path(project_dir)),
        # Contract: RMSD is mandatory for non-GNINA bridge paths as well.
        run_rmsd=True,
        run_visualizations=True,
        ligplus_root=ligplus_root,
        prompt_protein_names=prompt_protein_names,
        prompt_ligand_names=prompt_ligand_names,
        enable_poseview=enable_poseview,
        complex_query=complex_query,
        rmsd_scopes=rmsd_scopes,
        resume_rmsd=resume_rmsd,
        force_global_rmsd=force_global_rmsd,
        global_rmsd_defer_threshold=global_rmsd_defer_threshold,
        shared_mapping_dir=shared_mapping_dir,
        exclude_problematic_ligands=exclude_problematic_ligands,
        positive_affinity_threshold=positive_affinity_threshold,
        minimum_pose_count=minimum_pose_count,
        rmsd_workers=rmsd_workers,
        analysis_config=analysis_config,
    )
    if not bool(getattr(simplified, "run_rmsd", False)):
        raise RuntimeError("Non-GNINA analysis bridge contract violation: RMSD stage must be enabled")
    simplified.pairlist_df = analysis.pairlist_df.copy()
    simplified.receptor_files = find_receptor_files(shared_receptors_dir(project_dir))
    simplified.protein_name_map = build_protein_name_mapping(
        simplified.receptor_files,
        simplified.pairlist_df,
        overrides_file=simplified.protein_name_override_file,
    )
    simplified._prompt_for_protein_name_overrides()
    simplified._persist_protein_name_mapping()
    simplified.ligand_name_map = simplified._build_ligand_name_map()
    simplified._prompt_for_ligand_name_overrides()
    simplified._persist_ligand_name_mapping()
    simplified.scores_df = bridge_scores.copy()
    simplified.complexes = analysis._build_simplified_bridge_complexes(best_poses)

    complexes_dir = output_dir / "complexes"
    best_poses_pdb_dir = output_dir / "best_poses_pdb"
    manifest_df = best_poses.copy()
    manifest_df["bridge_output_name"] = manifest_df["tag"].astype(str)
    extracted_count = analysis._extract_best_pose_complexes(manifest_df, complexes_dir)
    analysis._mirror_complexes_to_best_poses(complexes_dir, best_poses_pdb_dir)

    return analysis, simplified, engine_scores, best_poses, extracted_count


def _run_non_gnina_analysis_target(
    target: str,
    project_dir: str,
    engine: str,
    output_dir: Optional[str] = None,
    ligplus_root: Optional[str] = None,
    enable_poseview: bool = False,
    prompt_protein_names: bool = False,
    prompt_ligand_names: bool = False,
    complex_query: Optional[str] = None,
    rmsd_scopes: Optional[object] = None,
    resume_rmsd: bool = True,
    force_global_rmsd: bool = False,
    global_rmsd_defer_threshold: int = 150,
    shared_mapping_dir: Optional[str] = None,
    exclude_problematic_ligands: bool = False,
    positive_affinity_threshold: float = 0.0,
    minimum_pose_count: int = 1,
    rmsd_workers: int = 0,
    speed_profile: str = "standard",
    analysis_config: Optional[Dict[str, object]] = None,
    analysis_config_file: str = "",
) -> WorkflowStepResult:
    if target in _BLOCKED_LEGACY_STAGE_TARGETS:
        _raise_legacy_target_migration_error(target)
    if target in _INTERACTION_TARGETS_ROUTED_TO_CLEAN:
        raise ValueError(
            f"Target {target} is an interaction alias. Dispatch it through run_analysis_target "
            "so it is normalized to analyze.interactions.clean and enforces the mandatory clean contract."
        )

    poseview_enabled = bool(enable_poseview or target == "analyze.interactions.poseview")
    root = _safe_root(project_dir=project_dir)
    out_dir = _resolve_analysis_output(Path(project_dir), output_dir, f"{engine}_{target.replace('.', '_')}")
    analysis, simplified, engine_scores, best_poses, extracted_count = _prepare_non_gnina_context(
        project_dir,
        engine,
        out_dir,
        ligplus_root,
        poseview_enabled,
        prompt_protein_names,
        prompt_ligand_names,
        complex_query,
        rmsd_scopes,
        resume_rmsd,
        force_global_rmsd,
        global_rmsd_defer_threshold,
        shared_mapping_dir,
        exclude_problematic_ligands,
        positive_affinity_threshold,
        minimum_pose_count,
        rmsd_workers,
        speed_profile,
        analysis_config,
    )

    success = extracted_count > 0
    if target == "analyze.stage.hierarchical":
        success = success and simplified._analyze_binding_affinity()
    elif target == "analyze.stage.polypharmacology":
        success = success and simplified._analyze_binding_affinity() and simplified._analyze_polypharmacology()
    elif target == "analyze.stage.rmsd":
        success = success and simplified._analyze_binding_affinity()
        if success:
            rmsd_result = analysis._run_non_gnina_rmsd_bridge(engine_scores, best_poses, out_dir / "complexes", out_dir)
            success = rmsd_result.get("state") == "completed"
    elif target == "analyze.stage.reports":
        success = success and simplified._generate_reports()
    elif target == "analyze.stage.visualizations":
        success = success and simplified._analyze_binding_affinity() and simplified._generate_visualizations()
    elif target == "analyze.interactions.poseview":
        success = success and simplified._analyze_binding_affinity() and simplified._generate_poseview_diagrams()
    elif target == "analyze.visuals.py3dmol":
        success = success and simplified._analyze_binding_affinity() and simplified._generate_py3dmol_visualizations()
    else:
        raise ValueError(f"Unsupported non-GNINA analysis target: {target}")

    if success:
        try:
            simplified._organize_output_artifacts()
        except Exception:
            pass

    _prune_empty_dirs(out_dir)
    update_context(
        root,
        project_root=str(Path(project_dir).expanduser().resolve()),
        analysis_output_dir=str(out_dir),
    )
    update_artifacts(root, last_analysis_output=str(out_dir))
    return _result(
        target,
        "completed" if success else "failed",
        root,
        inputs={
            "project_dir": project_dir,
            "engine": engine,
            "complex_query": complex_query or "",
            "exclude_problematic_ligands": exclude_problematic_ligands,
            "positive_affinity_threshold": positive_affinity_threshold,
            "minimum_pose_count": minimum_pose_count,
            "rmsd_workers": rmsd_workers,
            "speed_profile": speed_profile,
            "analysis_config_file": analysis_config_file,
        },
        outputs={"output_dir": str(out_dir)},
    )


def _extract_arg_value(argv: List[str], flag: str) -> Optional[str]:
    for idx, value in enumerate(argv):
        if value == flag and idx + 1 < len(argv):
            return argv[idx + 1]
        if value.startswith(f"{flag}="):
            return value.split("=", 1)[1]
    return None


def _resolve_analysis_output(root: Path, output_dir: Optional[str], target: str) -> Path:
    from docking.project_layout import detect_layout_profile, post_docking_root

    if output_dir:
        path = Path(output_dir).expanduser().resolve()
        path.mkdir(parents=True, exist_ok=True)
        return path
    target_slug = target.replace(".", "_")
    detect_layout_profile(root)
    path = post_docking_root(root) / "sessions" / target_slug
    path.mkdir(parents=True, exist_ok=True)
    return path


def _canonical_working_scores_root(root: Path) -> Path:
    from docking.project_layout import ensure_numbered_output_layout

    return ensure_numbered_output_layout(root)["post_scores_consensus"].parent


def _favorite_engine_from_manifest(project_dir: str) -> str:
    from docking.project_layout import manifest_path

    with open(manifest_path(Path(project_dir).expanduser().resolve()), "r", encoding="utf-8") as handle:
        payload = json.load(handle)
    favorite = str(payload.get("favorite_engine", "") or "")
    if favorite:
        return favorite
    engines = payload.get("engines", []) or []
    if len(engines) == 1:
        return str(engines[0])
    return ""
