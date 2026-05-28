from __future__ import annotations

from pathlib import Path

import pandas as pd


REQUIRED_SUMMARY_COLUMNS = {"PDB_ID", "Property", "Value"}
REQUIRED_PROPERTIES = {
    "active_site_center_x",
    "active_site_center_y",
    "active_site_center_z",
}


def load_site_catalog_from_summary(
    excel_path: Path,
    allow_missing_coordinates: bool = False,
    return_missing_coordinates: bool = False,
) -> pd.DataFrame | tuple[pd.DataFrame, list[str]]:
    excel_path = Path(excel_path)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read multi_pdb_analysis.xlsx") from exc

    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    if "Summary" not in workbook.sheetnames:
        raise ValueError(f"Summary sheet not found in {excel_path}")

    sheet = workbook["Summary"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Summary sheet is empty in {excel_path}")

    header = [str(value).strip() if value is not None else "" for value in rows[0]]
    frame = pd.DataFrame(rows[1:], columns=header)
    if not REQUIRED_SUMMARY_COLUMNS.issubset(frame.columns):
        missing = sorted(REQUIRED_SUMMARY_COLUMNS.difference(frame.columns))
        raise ValueError(f"Summary sheet missing required columns: {', '.join(missing)}")

    frame = frame[list(REQUIRED_SUMMARY_COLUMNS)].copy()
    frame["PDB_ID"] = frame["PDB_ID"].astype(str).str.upper().str.strip()
    frame["Property"] = frame["Property"].astype(str).str.strip()

    catalog = (
        frame.pivot_table(index="PDB_ID", columns="Property", values="Value", aggfunc="first")
        .reset_index()
        .rename_axis(None, axis=1)
    )
    rename_map = {
        "PDB_ID": "pdb_id",
        "selected_ligand": "selected_ligand",
        "protein_display_name": "protein_display_name",
        "active_site_center_x": "center_x",
        "active_site_center_y": "center_y",
        "active_site_center_z": "center_z",
    }
    for source, dest in rename_map.items():
        if source in catalog.columns:
            catalog[dest] = catalog[source]

    missing_props = sorted(prop for prop in REQUIRED_PROPERTIES if prop not in catalog.columns)
    if missing_props:
        raise ValueError(
            f"Summary sheet missing required properties: {', '.join(missing_props)}"
        )

    for axis in ("center_x", "center_y", "center_z"):
        catalog[axis] = pd.to_numeric(catalog[axis], errors="coerce")
    missing_coordinate_ids: list[str] = []
    if catalog[["center_x", "center_y", "center_z"]].isna().any().any():
        missing_coordinate_ids = sorted(
            catalog[catalog[["center_x", "center_y", "center_z"]].isna().any(axis=1)]["pdb_id"].astype(str).str.upper().tolist()
        )
        if not allow_missing_coordinates:
            raise ValueError(f"Missing active site coordinates for: {', '.join(missing_coordinate_ids)}")
        catalog = catalog[catalog[["center_x", "center_y", "center_z"]].notna().all(axis=1)].copy()

    if "selected_ligand" not in catalog.columns:
        catalog["selected_ligand"] = ""
    if "protein_display_name" not in catalog.columns:
        catalog["protein_display_name"] = ""

    result = catalog[
        ["pdb_id", "selected_ligand", "protein_display_name", "center_x", "center_y", "center_z"]
    ].copy()
    if return_missing_coordinates:
        return result, missing_coordinate_ids
    return result
