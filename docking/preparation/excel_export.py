"""Excel export helpers for PDB preparation results (Spec 038 R2).

EXCEL_AVAILABLE, excel_cell_value and the Summary-sheet writer moved here from
core_pipeline.py and from cli_pipeline.py / interactive_pipeline.py (where the
writer was duplicated). Behaviour is unchanged.
"""

import json

from docking.preparation.binding_site_extraction import centre_method_from_provenance


try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False


def excel_cell_value(value):
    """Serialise values openpyxl cannot write (dicts, lists) as JSON text (Spec 036 R6)."""
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, default=str, sort_keys=True)
    return value


def add_to_excel_workbook(workbook, pdb_id, results):
    """Upsert results into the Summary sheet by PDB ID."""
    try:
        # Create or get worksheet
        if "Summary" not in workbook.sheetnames:
            ws = workbook.create_sheet("Summary")
            # Add headers
            headers = ["PDB_ID", "Property", "Value"]
            ws.append(headers)
            
            # Style headers
            for cell in ws[1]:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center")
        else:
            ws = workbook["Summary"]

        normalized_pdb = str(pdb_id).strip().upper()

        # Replace prior entries for this PDB so reruns overwrite stale values.
        for row_idx in range(ws.max_row, 1, -1):
            cell_value = ws.cell(row=row_idx, column=1).value
            existing_pdb = str(cell_value).strip().upper() if cell_value is not None else ""
            if existing_pdb == normalized_pdb:
                ws.delete_rows(row_idx, 1)
        
        # Add results. Dict values are written as JSON text (Spec 036 R6).
        for key, value in results.items():
            ws.append([normalized_pdb, key, excel_cell_value(value)])
        centre_method = centre_method_from_provenance(results.get("binding_site_center_provenance"))
        if centre_method:
            ws.append([normalized_pdb, "binding_site_center_method", centre_method])
        
    except Exception as e:
        print(f"⚠️  Failed to add results to Excel: {e}")
